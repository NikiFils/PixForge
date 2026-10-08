#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
PixForge — универсальный конвертер изображений и PSD.

Возможности:
  • Пакетная обработка изображений с пошаговым превью
  • Входные форматы: PSD, PSB, TIFF, TGA, BMP, WebP, GIF, ICO,
                     PNG, JPEG, PDF, AI, HEIC/HEIF
  • Выходные форматы: PNG (прозр./бел.), JPG, WEBP, TIFF, BMP, PDF, AVIF
  • Фильтры по артикулам, папкам и датам
  • Импорт списка артикулов из TXT / CSV / XLSX / DOCX и Google Sheets
  • Автоматические отчёты по завершении
  • Тёмная / светлая / системная тема, адаптивная раскладка

Репозиторий:  https://github.com/NikiFils/PixForge
Автор:        NikiFils
Лицензия:     MIT
"""

import sys, os, io, re, csv, json, datetime, fnmatch, traceback, zipfile
import subprocess, importlib, importlib.util, threading, queue, time
import base64, xml.etree.ElementTree as ET
from pathlib import Path

# ============================================================================
#  МЕТАДАННЫЕ ПРИЛОЖЕНИЯ
# ============================================================================
APP_NAME    = "PixForge"
APP_VERSION = "1.0"
APP_TAGLINE = "Универсальный конвертер изображений и PSD"
APP_AUTHOR  = "NikiFils"
APP_YEAR    = "2026"
APP_LICENSE = "MIT"
APP_REPO    = "https://github.com/NikiFils/PixForge"

__version__ = APP_VERSION
__author__  = APP_AUTHOR
__license__ = APP_LICENSE

# ============================================================================
#  ПУТИ И ФАЙЛЫ
# ============================================================================
PIP_LOG  = Path(__file__).parent / "pixforge_pip.log"
LOG_PATH = Path(__file__).parent / "pixforge_error.log"
SETTINGS_FILE = Path.home() / ".pixforge_settings.json"
# Старые файлы настроек (для миграции с предыдущих версий)
_LEGACY_SETTINGS_FILES = [
    Path.home() / ".psd_converter_gui.json",
]

# ============================================================================
#  ЗАВИСИМОСТИ
# ============================================================================
REQUIRED_PACKAGES = [
    ("flet","flet","UI",True),("PIL","pillow","Изображения",True),
    ("psd_tools","psd-tools","PSD",True),
    ("openpyxl","openpyxl","Excel",False),
    ("gspread","gspread","Sheets",False),("google.auth","google-auth","Google",False),
    ("fitz","pymupdf","PDF/AI",False),
    ("pillow_heif","pillow-heif","HEIC",False),
    ("pillow_avif","pillow-avif-plugin","AVIF",False),
]

def _log_exc(t,v,tb):
    try:
        with LOG_PATH.open("a",encoding="utf-8") as f:
            f.write(f"\n--- {datetime.datetime.now():%Y-%m-%d %H:%M:%S} ---\n")
            traceback.print_exception(t,v,tb,file=f)
    except Exception:
        pass
sys.excepthook = _log_exc

def _is_installed(n):
    try:
        return importlib.util.find_spec(n) is not None
    except Exception:
        return False

def _is_frozen():
    return getattr(sys,"frozen",False)

def _real_python():
    e = sys.executable or "python"
    p = Path(e)
    if p.name.lower() == "pythonw.exe":
        c = p.with_name("python.exe")
        if c.is_file():
            return str(c)
    return e

def _pip_available():
    try:
        fl = subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0
        r = subprocess.run([_real_python(),"-m","pip","--version"],
                           capture_output=True,text=True,timeout=20,creationflags=fl)
        return r.returncode == 0
    except Exception:
        return False

def _missing():
    return [{"import_name":i,"pip_name":p,"description":d,"critical":c}
            for i,p,d,c in REQUIRED_PACKAGES if not _is_installed(i)]

def _restart():
    s = Path(__file__).resolve()
    try:
        args = ([sys.executable]+sys.argv[1:]) if _is_frozen() else [sys.executable,str(s)]+sys.argv[1:]
        fl = 0
        if sys.platform == "win32":
            fl = 0x00000008|0x00000200
            if Path(sys.executable).name.lower() == "pythonw.exe":
                fl |= 0x08000000
        subprocess.Popen(args,close_fds=False,creationflags=fl,cwd=str(Path(__file__).parent))
        return True
    except Exception:
        return False

def _install_gui():
    import tkinter as tk
    from tkinter import messagebox, scrolledtext
    miss = _missing()
    if not miss:
        return "ok"
    crit = [p for p in miss if p["critical"]]
    opt = [p for p in miss if not p["critical"]]
    L = []
    if crit:
        L.append("Обязательные:")
        for p in crit:
            L.append(f"  • {p['pip_name']} — {p['description']}")
    if opt:
        L.append("\nНеобязательные:")
        for p in opt:
            L.append(f"  • {p['pip_name']} — {p['description']}")
    if not _pip_available():
        r = tk.Tk(); r.withdraw()
        try:
            messagebox.showwarning("Зависимости","\n".join(L),parent=r)
        finally:
            r.destroy()
        return "ok" if not crit else "cancel"
    r = tk.Tk(); r.withdraw()
    try:
        if not messagebox.askyesno("Установка","\n".join(L+["\nУстановить?"]),parent=r,default="yes"):
            return "ok" if not crit else "cancel"
        win = tk.Toplevel(r); win.title("Установка"); win.geometry("820x580")
        win.transient(r); win.grab_set(); win.protocol("WM_DELETE_WINDOW",lambda:None)
        tk.Label(win,text="Устанавливаются:",font=("Segoe UI",11,"bold")).pack(anchor="w",padx=12,pady=(12,4))
        for p in miss:
            tk.Label(win,text=f"  • {p['pip_name']} — {p['description']}",fg="#333").pack(anchor="w",padx=12)
        logw = scrolledtext.ScrolledText(win,height=18,wrap="word",font=("Consolas",9))
        logw.pack(fill="both",expand=True,padx=12,pady=8)
        logw.configure(state="disabled")
        sv = tk.StringVar(value="Подготовка…")
        tk.Label(win,textvariable=sv,font=("Segoe UI",10,"bold")).pack(pady=(0,6))
        holder = {"ok":False}
        btn = tk.Button(win,text="Продолжить",state="disabled",font=("Segoe UI",11,"bold"),
                        padx=20,pady=8,command=win.destroy)
        btn.pack(pady=(0,12))
        q = queue.Queue()

        def alog(t):
            logw.configure(state="normal")
            logw.insert("end",t+"\n")
            logw.see("end")
            logw.configure(state="disabled")
            try:
                with PIP_LOG.open("a",encoding="utf-8") as f:
                    f.write(t+"\n")
            except Exception:
                pass

        def pip1(nm):
            cmd = [_real_python(),"-m","pip","install","--upgrade","--no-input",nm]
            fl = subprocess.CREATE_NO_WINDOW if sys.platform == "win32" else 0
            try:
                p = subprocess.Popen(cmd,stdout=subprocess.PIPE,stderr=subprocess.STDOUT,
                                     stdin=subprocess.DEVNULL,text=True,encoding="utf-8",
                                     errors="replace",creationflags=fl,bufsize=1)
                for ln in p.stdout:
                    q.put(("log",ln.rstrip()))
                p.wait()
                return p.returncode == 0
            except Exception as ex:
                q.put(("log",f"Ошибка: {ex}"))
                return False

        def work():
            ok = True
            tot = len(miss)
            for i,p in enumerate(miss,1):
                q.put(("status",f"[{i}/{tot}] {p['pip_name']}…"))
                q.put(("log",f"\n=== pip install {p['pip_name']} ==="))
                importlib.invalidate_caches()
                r1 = pip1(p["pip_name"])
                q.put(("log",f"[{'OK' if r1 and _is_installed(p['import_name']) else 'FAIL'}]"))
                if not (r1 and _is_installed(p["import_name"])):
                    ok = False
            q.put(("done",ok))

        def poll():
            try:
                while True:
                    m = q.get_nowait()
                    if m[0] == "log":
                        alog(m[1])
                    elif m[0] == "status":
                        sv.set(m[1])
                    elif m[0] == "done":
                        holder["ok"] = m[1]
                        if m[1]:
                            sv.set("✔ Готово.")
                            btn.config(state="normal",bg="#16a34a",fg="white")
                            win.after(2000,win.destroy)
                        else:
                            sv.set("⚠ Ошибки.")
                            btn.config(state="normal",bg="#dc2626",fg="white")
                        win.protocol("WM_DELETE_WINDOW",win.destroy)
                        return
            except queue.Empty:
                pass
            win.after(80,poll)

        threading.Thread(target=work,daemon=True).start()
        win.after(60,poll)
        r.wait_window(win)
        if holder["ok"]:
            return "restart"
        if any(p["critical"] for p in _missing()):
            messagebox.showerror("Не удалось","Обязательные не установились.",parent=r)
            return "cancel"
        return "restart"
    finally:
        try:
            r.destroy()
        except Exception:
            pass

def _maybe_restart():
    st = _install_gui()
    if st == "restart":
        if _restart():
            raise SystemExit(0)
        raise SystemExit(0)
    elif st == "cancel":
        raise SystemExit(0)

_maybe_restart()

import flet as ft
from PIL import Image
from psd_tools import PSDImage

# ============================================================================
#  КОНСТАНТЫ ИНТЕРФЕЙСА
# ============================================================================
GAP        = 10
GAP_FIELD  = 6
PAD        = 12
ROW_H      = 40
ROW_H_MAIN = 48
BTN_RADIUS = 8
LBL_SIZE   = 12
LBL_WEIGHT = ft.FontWeight.W_600
TXT_SIZE   = 11
HINT_SIZE  = 10
SPACER     = 10

LEFT_MIN       = 260
RIGHT_MIN      = 280
SASH_W         = 10
WIN_MIN_WIDTH  = 640
WIN_MIN_HEIGHT = 500

NARROW_W      = 420
VERY_NARROW_W = 320
RIGHT_NARROW  = 420

def _safe(cls, **kw):
    for _ in range(30):
        try:
            return cls(**kw)
        except TypeError as e:
            msg = str(e)
            m = re.search(r"unexpected keyword argument '(\w+)'", msg)
            if m and m.group(1) in kw:
                kw.pop(m.group(1))
                continue
            names = re.findall(r"'(\w+)'", msg)
            rem = False
            for k in names:
                if k in kw:
                    kw.pop(k)
                    rem = True
                    break
            if rem:
                continue
            raise

def _dd(**kw):
    cb = kw.pop("on_change",None) or kw.pop("on_select",None)
    kw.setdefault("text_size", TXT_SIZE)
    kw.setdefault("border_radius", BTN_RADIUS)
    try:
        kw.setdefault("text_style", ft.TextStyle(size=TXT_SIZE, weight=ft.FontWeight.NORMAL))
    except Exception:
        pass
    try:
        kw.setdefault("hint_style", ft.TextStyle(size=TXT_SIZE, weight=ft.FontWeight.NORMAL))
    except Exception:
        pass
    try:
        kw.setdefault("bgcolor", C_SURF_HIGH)
    except Exception:
        pass
    try:
        kw.setdefault("fill_color", C_SURF_HIGHEST)
    except Exception:
        pass
    if cb is not None:
        for key in ("on_select","on_change"):
            try:
                return _safe(ft.Dropdown, **kw, **{key:cb})
            except TypeError:
                continue
    return _safe(ft.Dropdown, **kw)

def _cb(**kw):
    kw.setdefault("checkbox_width", 18)
    kw.setdefault("checkbox_height", 18)
    return _safe(ft.Checkbox, **kw)

def _tf(**kw):
    kw.setdefault("text_size", TXT_SIZE)
    kw.setdefault("border_radius", BTN_RADIUS)
    kw.setdefault("min_lines", 1)
    try:
        kw.setdefault("text_style", ft.TextStyle(size=TXT_SIZE, weight=ft.FontWeight.NORMAL))
    except Exception:
        pass
    return _safe(ft.TextField, **kw)

def _tf_multi(**kw):
    kw.setdefault("text_size", TXT_SIZE)
    kw.setdefault("border_radius", BTN_RADIUS)
    kw.setdefault("multiline", True)
    kw.setdefault("min_lines", 1)
    kw.setdefault("max_lines", 6)
    try:
        kw.setdefault("text_style", ft.TextStyle(size=TXT_SIZE, weight=ft.FontWeight.NORMAL))
    except Exception:
        pass
    return _safe(ft.TextField, **kw)

def _btn_style(**kw):
    r = kw.pop("_radius", BTN_RADIUS)
    b = getattr(ft,"RoundedRectangleBorder",None)
    if b:
        try:
            kw["shape"] = b(radius=r)
        except Exception:
            pass
    return _safe(ft.ButtonStyle, **kw)

def _ev(e, default=None):
    if e is None:
        return default
    c = getattr(e,"control",None)
    if c is not None:
        v = getattr(c,"value",None)
        if v is not None:
            return v
    d = getattr(e,"data",None)
    return d if d is not None else default

def _pad_only(**kw):
    for cls_name in ("Padding","padding"):
        cls = getattr(ft, cls_name, None)
        if cls is None:
            continue
        fn = getattr(cls, "only", None)
        if fn is None:
            continue
        try:
            return fn(**kw)
        except Exception:
            pass
    vals = [v for v in kw.values() if isinstance(v,(int,float))]
    return max(vals) if vals else 0

def _pad_sym(h=0, v=0):
    for cls_name in ("Padding","padding"):
        cls = getattr(ft, cls_name, None)
        if cls is None:
            continue
        fn = getattr(cls, "symmetric", None)
        if fn is None:
            continue
        try:
            return fn(horizontal=h, vertical=v)
        except Exception:
            pass
    return max(h, v)

def _align(name):
    for n in ("Alignment","alignment"):
        c = getattr(ft,n,None)
        if c is None:
            continue
        a = getattr(c,name.upper(),None) or getattr(c,name.lower(),None)
        if a:
            return a

CENTER      = _align("center")
TOP_CENTER  = _align("top_center")
TOP_LEFT    = _align("top_left")

def _fit(name):
    for n in ("BoxFit","ImageFit"):
        c = getattr(ft,n,None)
        if c is None:
            continue
        a = getattr(c,name.upper(),None) or getattr(c,name.lower(),None)
        if a:
            return a

BOX_CONTAIN = _fit("contain")

def _border_all(w, c):
    for cls_name in ("Border","border"):
        cls = getattr(ft, cls_name, None)
        if cls is None:
            continue
        fn = getattr(cls, "all", None)
        if fn is None:
            continue
        try:
            return fn(w, c)
        except Exception:
            pass
    return None

def _col(name, fallback):
    try:
        c = getattr(ft.Colors,name,None)
        if c is not None:
            return c
    except Exception:
        pass
    return fallback

C_SURF_HIGH    = _col("SURFACE_CONTAINER_HIGH", _col("SURFACE_VARIANT","#e8e8e8"))
C_SURF_HIGHEST = _col("SURFACE_CONTAINER_HIGHEST", _col("SURFACE","#f0f0f0"))
C_ON_VAR       = _col("ON_SURFACE_VARIANT", _col("ON_SURFACE","#888"))
C_ON_SURF      = _col("ON_SURFACE", "#e8e8e8")
C_OUT_VAR      = _col("OUTLINE_VARIANT", _col("OUTLINE","#ccc"))
BG_APP         = _col("SURFACE", "#f7f7f7")

def _run_thread(page, fn):
    try:
        page.run_thread(fn)
        return
    except Exception:
        pass
    try:
        page.run_task(fn)
        return
    except Exception:
        pass
    threading.Thread(target=fn,daemon=True).start()

# ============================================================================
#  ФОРМАТЫ ФАЙЛОВ
# ============================================================================
INPUT_EXT_INFO = {
    ".psd":"Photoshop",".psb":"Photoshop (большой)",
    ".tif":"TIFF",".tiff":"TIFF",".tga":"TGA",".bmp":"BMP",
    ".webp":"WebP",".gif":"GIF",".ico":"ICO",
    ".png":"PNG",".jpg":"JPEG",".jpeg":"JPEG",
    ".pdf":"PDF",".ai":"Adobe Illustrator",
    ".heic":"HEIC",".heif":"HEIF",
}
INPUT_PRESETS = {
    "Только PSD":[".psd",".psb"],
    "PSD + растровые":[".psd",".psb",".tif",".tiff",".tga",".bmp",
                       ".webp",".gif",".ico",".png",".jpg",".jpeg"],
    "Всё поддерживаемое":list(INPUT_EXT_INFO.keys()),
}
INPUT_PRESET_CUSTOM = "Свой"
PRESET_ICONS = {
    "Только PSD": "PHOTO_SIZE_SELECT_ACTUAL",
    "PSD + растровые": "COLLECTIONS",
    "Всё поддерживаемое": "ALL_INCLUSIVE",
}
OUTPUT_FORMATS = [
    ("png_t","PNG-прозр",True),("png_o","PNG-бел",False),
    ("jpg","JPG",False),("webp","WEBP",False),("tiff","TIFF",False),
    ("bmp","BMP",False),("pdf","PDF",False),("avif","AVIF",False),
]
OUTPUT_REQUIRED_LIB = {"avif":("pillow_avif","pillow-avif-plugin")}

# ============================================================================
#  НАСТРОЙКИ ПО УМОЛЧАНИЮ
# ============================================================================
DEFAULT_SETTINGS = {
    "theme":"system","left_panel_width":520,"dark_preview":False,
    "input_exts":list(INPUT_PRESETS["Только PSD"]),
    "overwrite_mode":"overwrite","match_mode":"partial",
    "articles":"","include_folders":"","exclude_folders":"",
    "date_mode":"all","date_kind":"modified",
    "date_from":"","date_to":"",
    "last_folder":"","save_stats":False,"report_dir":"",
    "last_xlsx_column":"","gsheet_url":"","gsheet_method":"public",
    "gsheet_creds":"","gsheet_column":"",
    "sec_open_formats":True,"sec_open_articles":True,
    "sec_open_filters":True,"sec_open_stats":True,
    "show_tooltips":True,
    **{f"fmt_{fid}":d for fid,_,d in OUTPUT_FORMATS},
}

def load_settings():
    # Автомиграция со старых файлов настроек.
    if not SETTINGS_FILE.exists():
        for legacy in _LEGACY_SETTINGS_FILES:
            try:
                if legacy.exists():
                    SETTINGS_FILE.write_text(
                        legacy.read_text(encoding="utf-8"),
                        encoding="utf-8")
                    break
            except Exception:
                continue
    try:
        d = json.loads(SETTINGS_FILE.read_text(encoding="utf-8"))
        s = DEFAULT_SETTINGS.copy()
        s.update(d)
        return s
    except Exception:
        return DEFAULT_SETTINGS.copy()

def save_settings(s):
    try:
        SETTINGS_FILE.write_text(json.dumps(s,ensure_ascii=False,indent=2),encoding="utf-8")
    except Exception:
        pass

# ============================================================================
#  ВСПОМОГАТЕЛЬНЫЕ ФУНКЦИИ
# ============================================================================
def parse_lines(t):
    if not t:
        return []
    parts = re.split(r"[,\n\r;]+", t)
    return [p.strip() for p in parts if p and p.strip()]

def app_dir():
    if getattr(sys,"frozen",False):
        return Path(sys.executable).parent
    return Path(__file__).parent

def scan_files(root, exts, progress_cb=None, cancel_check=None):
    paths, meta, stack, n = [], {}, [root], 0
    while stack:
        if cancel_check and cancel_check():
            break
        cur = stack.pop()
        try:
            with os.scandir(cur) as it:
                for e in it:
                    try:
                        if e.is_dir(follow_symlinks=False):
                            stack.append(Path(e.path))
                        elif e.is_file(follow_symlinks=False):
                            if os.path.splitext(e.name)[1].lower() in exts:
                                p = Path(e.path)
                                try:
                                    st = e.stat(follow_symlinks=False)
                                    meta[p] = {"mtime":st.st_mtime,"ctime":st.st_ctime}
                                except OSError:
                                    meta[p] = {"mtime":0.0,"ctime":0.0}
                                paths.append(p)
                    except OSError:
                        continue
        except OSError:
            continue
        n += 1
        if progress_cb and n % 30 == 0:
            progress_cb(n,len(paths))
    return paths, meta

def load_image_any(path, pdf_page=0):
    ext = path.suffix.lower()
    if ext in (".psd",".psb"):
        psd = PSDImage.open(path)
        img = psd.composite()
        if img is None:
            raise ValueError("Пустой PSD")
        return img.convert("RGBA")
    if ext in (".pdf",".ai"):
        try:
            import fitz
        except ImportError:
            try:
                import pymupdf as fitz
            except Exception:
                raise RuntimeError("Нужна pymupdf")
        doc = fitz.open(str(path))
        try:
            if len(doc) == 0:
                raise ValueError("Пустой документ")
            idx = min(max(0,pdf_page),len(doc)-1)
            pix = doc.load_page(idx).get_pixmap(alpha=True)
            return Image.frombytes("RGBA",(pix.width,pix.height),pix.samples)
        finally:
            doc.close()
    if ext in (".heic",".heif"):
        try:
            from pillow_heif import register_heif_opener
            register_heif_opener()
        except ImportError:
            raise RuntimeError("Нужна pillow-heif")
    img = Image.open(path)
    try:
        img.seek(0)
    except Exception:
        pass
    return img.convert("RGBA")

def save_image_in_format(img,out_path,kind):
    if kind == "png_t":
        img.save(out_path,"PNG")
    elif kind == "png_o":
        bg = Image.new("RGB",img.size,(255,255,255))
        m = img.split()[-1] if img.mode == "RGBA" else None
        bg.paste(img,mask=m)
        bg.save(out_path,"PNG")
    elif kind == "jpg":
        if img.mode == "RGBA":
            bg = Image.new("RGB",img.size,(255,255,255))
            bg.paste(img,mask=img.split()[-1])
            j = bg
        else:
            j = img.convert("RGB")
        j.save(out_path,"JPEG",quality=92)
    elif kind == "webp":
        img.save(out_path,"WEBP",quality=92,method=6)
    elif kind == "tiff":
        img.save(out_path,"TIFF",compression="tiff_lzw")
    elif kind == "bmp":
        if img.mode == "RGBA":
            bg = Image.new("RGB",img.size,(255,255,255))
            bg.paste(img,mask=img.split()[-1])
            b = bg
        else:
            b = img.convert("RGB")
        b.save(out_path,"BMP")
    elif kind == "pdf":
        if img.mode != "RGBA":
            img = img.convert("RGBA")
        img.save(out_path,"PDF",resolution=150.0)
    elif kind == "avif":
        try:
            import pillow_avif  # noqa
        except ImportError:
            raise RuntimeError("pip install pillow-avif-plugin")
        img.save(out_path,"AVIF",quality=80)

def pil_to_data_uri(img,mw=1400,mh=1400):
    c = img.copy()
    c.thumbnail((mw,mh),Image.LANCZOS)
    if c.mode not in ("RGB","RGBA"):
        c = c.convert("RGBA")
    buf = io.BytesIO()
    c.save(buf,format="PNG",optimize=True)
    b64 = base64.b64encode(buf.getvalue()).decode()
    return f"data:image/png;base64,{b64}"

def _wildcard(pat):
    if not pat:
        return pat
    if "*" in pat or "?" in pat:
        return pat.lower()
    return f"*{pat.lower()}*"

# ============================================================================
#  ЧТЕНИЕ СПИСКОВ ИЗ ФАЙЛОВ
# ============================================================================
def read_xlsx(path):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RuntimeError("Нужен openpyxl")
    wb = load_workbook(filename=str(path),read_only=True,data_only=True)
    ws = wb.active
    cols = {}
    for row in ws.iter_rows(values_only=False):
        for c in row:
            if c.value is None:
                continue
            v = str(c.value).strip()
            if not v:
                continue
            cols.setdefault(c.column_letter,[]).append(v)
    wb.close()
    o = sorted(cols.keys(),key=lambda x:(len(x),x))
    return o,{k:cols[k] for k in o}

def read_docx(path):
    try:
        with zipfile.ZipFile(path,"r") as z:
            if "word/document.xml" not in z.namelist():
                raise ValueError("нет document.xml")
            with z.open("word/document.xml") as f:
                tree = ET.parse(f)
    except Exception as e:
        raise RuntimeError(f"docx: {e}")
    lines = []
    for t in tree.getroot().iter("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"):
        if t.text and t.text.strip():
            lines.append(t.text.strip())
    return lines

# ============================================================================
#  GOOGLE SHEETS
# ============================================================================
def parse_gsheet_url(url):
    m = re.search(r"/spreadsheets/d/([a-zA-Z0-9\-_]+)",url.strip())
    if not m:
        raise ValueError("Не найден ID")
    sid = m.group(1)
    gm = re.search(r"[#&?]gid=(\d+)",url)
    gid = int(gm.group(1)) if gm else None
    return sid,gid

def _csv_letter(i):
    r,i = "",i+1
    while i > 0:
        i,rem = divmod(i-1,26)
        r = chr(65+rem)+r
    return r

def _rows_to_columns(rows):
    if not rows:
        return [],{}
    mc = max(len(r) for r in rows)
    cols = {}
    for ci in range(mc):
        letter = _csv_letter(ci)
        vals = []
        for r in rows:
            if ci < len(r):
                v = str(r[ci]).strip()
                if v:
                    vals.append(v)
        if vals:
            cols[letter] = vals
    o = sorted(cols.keys(),key=lambda x:(len(x),x))
    return o,{k:cols[k] for k in o}

def read_gsheet_public(url):
    sid,gid = parse_gsheet_url(url)
    u = f"https://docs.google.com/spreadsheets/d/{sid}/export?format=csv"
    if gid is not None:
        u += f"&gid={gid}"
    import urllib.request, urllib.error
    try:
        with urllib.request.urlopen(urllib.request.Request(u,headers={"User-Agent":"Mozilla/5.0"}),timeout=30) as r:
            raw = r.read()
    except urllib.error.HTTPError as e:
        if e.code in (401,403):
            raise RuntimeError("Таблица не публичная")
        raise RuntimeError(f"HTTP {e.code}")
    except urllib.error.URLError as e:
        raise RuntimeError(f"Сеть: {e.reason}")
    return _rows_to_columns(list(csv.reader(io.StringIO(raw.decode("utf-8",errors="replace")))))

def read_gsheet_private(url, creds_path):
    try:
        import gspread
        from google.oauth2.service_account import Credentials
    except ImportError:
        raise RuntimeError("pip install gspread google-auth")
    cf = Path(creds_path)
    if not cf.is_file():
        raise RuntimeError(f"Файл ключа не найден: {cf}")
    scopes = ["https://www.googleapis.com/auth/spreadsheets.readonly",
              "https://www.googleapis.com/auth/drive.readonly"]
    creds = Credentials.from_service_account_file(str(cf),scopes=scopes)
    gc = gspread.authorize(creds)
    sid,gid = parse_gsheet_url(url)
    sh = gc.open_by_key(sid)
    try:
        ws = sh.get_worksheet_by_id(gid) if gid is not None else sh.get_worksheet(0)
    except Exception:
        ws = sh.get_worksheet(0)
    if ws is None:
        raise RuntimeError("Лист не найден")
    return _rows_to_columns(ws.get_all_values())

# ============================================================================
#  UI-КОМПОНЕНТЫ
# ============================================================================
class Section(ft.Container):
    def __init__(self, title, body_controls, key, app, initially_open=True):
        super().__init__()
        self._key = key
        self._app = app
        self._open = bool(initially_open)
        self._title_text = ft.Text(title, size=13, weight=ft.FontWeight.BOLD,
                                    no_wrap=True, overflow=ft.TextOverflow.ELLIPSIS,
                                    expand=True)
        self._summary_text = ft.Text("", size=HINT_SIZE, color=C_ON_VAR, italic=True,
                                       no_wrap=True, overflow=ft.TextOverflow.ELLIPSIS)
        self._body = ft.Container(content=ft.Column(body_controls, spacing=GAP_FIELD, tight=True),
                                   padding=PAD, visible=self._open)
        self._chevron = ft.Text("▾" if self._open else "▸",
                                 size=14, weight=ft.FontWeight.BOLD, color=ft.Colors.PRIMARY)
        self._header = ft.Container(
            content=ft.Row([self._chevron,
                self._title_text,
                self._summary_text], spacing=GAP,
                vertical_alignment=ft.CrossAxisAlignment.CENTER),
            padding=_pad_sym(PAD, GAP), on_click=self._toggle, border_radius=12)
        self.bgcolor = C_SURF_HIGH
        self.border_radius = 14
        self.content = ft.Column([self._header, self._body], spacing=0, tight=True)

    def _toggle(self, e):
        self._open = not self._open
        self._body.visible = self._open
        self._chevron.value = "▾" if self._open else "▸"
        try:
            self._app.settings[self._key] = self._open
            save_settings(self._app.settings)
        except Exception:
            pass
        try:
            self.update()
        except Exception:
            pass

    def set_summary(self, text):
        self._summary_text.value = text or ""

    def set_compact(self, compact):
        try:
            self._summary_text.visible = not compact
            self._summary_text.update()
        except Exception:
            pass

    def set_title(self, text):
        try:
            self._title_text.value = text
            self._title_text.update()
        except Exception:
            pass

    def set_body_padding(self, value):
        try:
            self._body.padding = value
            self._body.update()
        except Exception:
            pass

class ThemeToggle(ft.Container):
    def __init__(self, value=False, on_change=None, **kwargs):
        self.value = bool(value)
        self._cb = on_change
        self._track_w = 48
        self._track_h = 26
        self._thumb_size = 20
        self._pad = 3
        self._thumb = ft.Container(
            width=self._thumb_size, height=self._thumb_size,
            border_radius=self._thumb_size // 2,
            bgcolor=ft.Colors.WHITE,
            left=self._pad if not self.value else (self._track_w - self._thumb_size - self._pad),
            top=self._pad,
        )
        self._track = ft.Container(
            width=self._track_w, height=self._track_h,
            border_radius=self._track_h // 2,
            bgcolor=ft.Colors.PRIMARY if self.value else C_SURF_HIGHEST,
            content=ft.Stack([self._thumb], width=self._track_w, height=self._track_h),
            on_click=self._toggle, ink=False,
        )
        super().__init__(content=self._track, **kwargs)

    def _toggle(self, e):
        self.value = not self.value
        new_left = self._pad if not self.value else (self._track_w - self._thumb_size - self._pad)
        try:
            self._thumb.left = new_left
            self._track.bgcolor = ft.Colors.PRIMARY if self.value else C_SURF_HIGHEST
            self._thumb.update()
            self._track.update()
        except Exception:
            pass
        if self._cb:
            try:
                class _FakeEv:
                    pass
                fe = _FakeEv()
                fe.control = self
                fe.data = str(self.value).lower()
                self._cb(fe)
            except Exception:
                pass

# ============================================================================
#  ОСНОВНОЕ ПРИЛОЖЕНИЕ
# ============================================================================
class ConverterApp:
    def __init__(self, page: ft.Page):
        self.page = page
        page.title = APP_NAME
        try:
            page.window.width = 1280
            page.window.height = 900
            page.window.min_width = WIN_MIN_WIDTH
            page.window.min_height = WIN_MIN_HEIGHT
        except Exception:
            pass
        page.padding = PAD
        page.spacing = 0
        self.settings = load_settings()
        self._apply_theme(self.settings.get("theme","system"))

        self.picker = ft.FilePicker()
        for attr in ("services","overlay"):
            lst = getattr(page,attr,None)
            if lst is None:
                continue
            try:
                lst.append(self.picker)
                break
            except Exception:
                continue

        self.files = []
        self.idx = 0
        self.converted = 0
        self.skipped = 0
        self.errors = 0
        self.full_image = None
        self.current_folder = None
        self.all_found_count = 0
        self.all_after_folder_filter = 0
        self.all_after_date_filter = 0
        self.finished = True
        self.log_entries = []
        self.processed = set()
        self._file_cache = None
        self._file_meta = {}
        self._scan_should_stop = False
        self._scanning = False
        self._left_panel_width = self.settings.get("left_panel_width",520)
        self._input_exts = set(self.settings.get("input_exts",INPUT_PRESETS["Только PSD"]))
        self.out_vars = {fid:self.settings.get(f"fmt_{fid}",d) for fid,_,d in OUTPUT_FORMATS}
        self._debounce_timer = None
        self._sash_start_w = 520
        self._sash_last_gx = None
        self._left_scroll_max = 0.0
        self._left_scroll_pos = 0.0
        self._scroll_drag_start = None
        self._sash_mouse_start_x = None
        self._out_containers = []
        self._folder_pair_stacked = None
        self._resize_timer = None
        self._resize_event_width = None
        self._last_applied_width = None
        self._watcher_stop = False
        self._active_dialog = None
        self._tooltip_targets = []
        self._top_row_alignment = None

        self._build_ui()
        try:
            page.on_keyboard_event = self._on_key
        except Exception:
            pass
        try:
            page.on_resize = self._on_resize
        except Exception:
            pass
        self._refresh_section_summaries()
        self._register_tooltips()
        self._apply_tooltips()
        page.update()

        win_w = self._read_page_width_safe()
        self._apply_width(win_w, force=True)

        threading.Thread(target=self._width_watcher, daemon=True).start()

    # ---------- ЧТЕНИЕ ШИРИНЫ СТРАНИЦЫ ----------
    def _read_page_width_safe(self):
        candidates = []
        try:
            v = getattr(self.page, "width", None)
            if v and float(v) > 100:
                return float(v)
        except Exception:
            pass
        try:
            v = self.page.window.width
            if v and float(v) > 100:
                candidates.append(float(v))
        except Exception:
            pass
        if candidates:
            return min(candidates)
        return 1280.0

    @staticmethod
    def _extract_width_from_event(e):
        if e is None:
            return None
        for attr in ("width","w","page_width","client_width"):
            v = getattr(e, attr, None)
            if v is not None:
                try:
                    f = float(v)
                    if f > 100:
                        return f
                except Exception:
                    pass
        data = getattr(e, "data", None)
        if data is None:
            return None
        try:
            if isinstance(data, dict):
                for k in ("width","w","page_width"):
                    if data.get(k):
                        return float(data[k])
        except Exception:
            pass
        s = str(data).strip()
        if not s:
            return None
        try:
            d = json.loads(s)
            if isinstance(d, dict):
                for k in ("width","w","page_width"):
                    if d.get(k):
                        return float(d[k])
            elif isinstance(d, list) and d:
                return float(d[0])
        except Exception:
            pass
        for sep in (",","x",";"," "):
            if sep in s:
                try:
                    w = float(s.split(sep)[0].strip())
                    if w > 100:
                        return w
                except Exception:
                    pass
        return None

    def _apply_width(self, w, force=False):
        try:
            w = float(w)
            if w < 200:
                return
            if self._last_applied_width is not None and not force:
                if w >= self._last_applied_width - 1 and abs(w - self._last_applied_width) < 2:
                    return
            self._last_applied_width = w

            max_left = max(LEFT_MIN, w - SASH_W - RIGHT_MIN)
            changed = False
            if self._left_panel_width > max_left:
                self._left_panel_width = max_left
                try:
                    self._left_box.width = max_left
                except Exception:
                    pass
                self.settings["left_panel_width"] = max_left
                changed = True
            if force or changed:
                try:
                    self._left_box.update()
                except Exception:
                    pass
            self._apply_adaptive(self._left_panel_width)
            self._apply_right_adaptive(w - self._left_panel_width - SASH_W)
        except Exception:
            pass

    def _width_watcher(self):
        while not self._watcher_stop:
            time.sleep(0.4)
            try:
                w = self._read_page_width_safe()
                self._apply_width(w)
            except Exception:
                pass

    def _apply_theme(self, mode):
        try:
            if mode == "dark":
                self.page.theme_mode = ft.ThemeMode.DARK
            elif mode == "light":
                self.page.theme_mode = ft.ThemeMode.LIGHT
            else:
                self.page.theme_mode = ft.ThemeMode.SYSTEM
        except Exception:
            pass

    def _on_theme_change(self, e):
        v = _ev(e, "system")
        mp = {"light":"light","dark":"dark","system":"system",
              "Светлая":"light","Тёмная":"dark","Системная":"system"}
        mode = mp.get(v, "system")
        self.settings["theme"] = mode
        save_settings(self.settings)
        self._apply_theme(mode)
        try:
            self.page.update()
        except Exception:
            pass

    def _snack(self, msg):
        try:
            sb = ft.SnackBar(ft.Text(msg), duration=4000)
            if sb not in self.page.overlay:
                self.page.overlay.append(sb)
            sb.open = True
            self.page.update()
        except Exception:
            print(msg)

    def _open_dlg(self, dlg):
        self._active_dialog = dlg
        try:
            self.page.run_task(self.page.open, dlg)
            return
        except Exception:
            pass
        try:
            if dlg not in self.page.overlay:
                self.page.overlay.append(dlg)
            dlg.open = True
            self.page.update()
        except Exception:
            pass

    def _close_dlg(self, dlg):
        if self._active_dialog is dlg:
            self._active_dialog = None
        try:
            self.page.run_task(self.page.close, dlg)
            return
        except Exception:
            pass
        try:
            dlg.open = False
            self.page.update()
        except Exception:
            pass

    def _card(self,*ctrls,spacing=GAP,pad=PAD):
        return ft.Container(content=ft.Column(list(ctrls),spacing=spacing,tight=True),
                             padding=pad, bgcolor=C_SURF_HIGH, border_radius=14)

    def _chip(self, icon_name, text, on_click, selected=False):
        icon = getattr(ft.Icons, icon_name, None) or ft.Icons.CIRCLE
        bg = ft.Colors.PRIMARY if selected else C_SURF_HIGHEST
        fg = ft.Colors.WHITE if selected else C_ON_VAR
        return ft.Container(
            content=ft.Row([
                ft.Icon(icon, size=14, color=fg),
                ft.Text(text, size=TXT_SIZE, weight=ft.FontWeight.W_500, color=fg),
            ], spacing=5, tight=True, alignment=ft.MainAxisAlignment.CENTER),
            padding=_pad_sym(11, 7), border_radius=16,
            bgcolor=bg, on_click=on_click, ink=True,
        )

    def _kbd(self, key):
        return ft.Container(
            content=ft.Text(key, size=HINT_SIZE, weight=ft.FontWeight.W_600,
                             color=C_ON_SURF, font_family="Consolas"),
            padding=_pad_sym(7, 3),
            bgcolor=C_SURF_HIGH, border_radius=5,
        )

    def _hk(self, key, desc):
        return ft.Row([self._kbd(key), ft.Text(desc, size=HINT_SIZE, color=C_ON_VAR)],
                       spacing=4, tight=True)

    def _field(self, label, control, hint=None):
        ctrls = []
        if label:
            ctrls.append(ft.Text(label, size=LBL_SIZE, weight=LBL_WEIGHT, color=C_ON_SURF,
                                   no_wrap=False))
        ctrls.append(control)
        if hint:
            ctrls.append(ft.Container(content=ft.Text(hint, size=HINT_SIZE,
                                                       color=C_ON_VAR, italic=True,
                                                       no_wrap=False),
                                       padding=_pad_only(top=2)))
        return ft.Column(ctrls, spacing=GAP_FIELD, tight=True)

    def _spacer(self, h=SPACER):
        return ft.Container(height=h)

    def _set_folder_pair_stacked(self, stacked):
        if stacked == self._folder_pair_stacked:
            return
        self._folder_pair_stacked = stacked
        try:
            if stacked:
                self.folder_path.expand = None
                self.btn_choose_folder.expand = None
                self._folder_pair.content = ft.Column(
                    [self.btn_choose_folder, self.folder_path],
                    spacing=GAP_FIELD, tight=True,
                    horizontal_alignment=ft.CrossAxisAlignment.STRETCH,
                )
            else:
                self.folder_path.expand = True
                self.btn_choose_folder.expand = None
                self._folder_pair.content = ft.Row(
                    [self.btn_choose_folder, self.folder_path],
                    spacing=GAP,
                    vertical_alignment=ft.CrossAxisAlignment.CENTER,
                )
            self._folder_pair.update()
        except Exception:
            pass

    def _set_top_row_alignment(self, alignment):
        if alignment == self._top_row_alignment:
            return
        self._top_row_alignment = alignment
        try:
            self._top_row.alignment = alignment
            self._top_row.update()
        except Exception:
            pass

    # ---------- TOOLTIPS ----------
    def _register_tooltips(self):
        items = [
            (self.folder_path,         "Путь к папке с исходниками (можно вставить вручную)"),
            (self.btn_choose_folder,   "Выбрать папку с исходными файлами"),
            (self.btn_rescan,          "Пересканировать текущую папку"),
            (self.btn_cancel_scan,     "Прервать сканирование"),
            (self.preset_dd,           "Набор входных расширений"),
            (self.btn_setup,           "Настроить список входных форматов"),
            (self.btn_settings,        "Общие настройки программы"),
            (self.btn_help,            "Открыть справку"),
            (self.btn_about,           "О программе"),
            (self.overwrite_dd,        "Что делать, если выходной файл уже существует"),
            (self.match_dd,            "Как сравнивать артикулы с именами файлов"),
            (self.btn_from_file,       "Загрузить список артикулов из файла"),
            (self.btn_from_gsheet,     "Загрузить артикулы из Google Sheets"),
            (self.btn_clear_search,    "Очистить список артикулов"),
            (self.btn_apply_filter,    "Применить фильтр артикулов сейчас"),
            (self.date_dd,             "Период по дате файла"),
            (self.date_kind_dd,        "Что считать датой: изменение или создание"),
            (self.date_from_field,     "Начало периода"),
            (self.date_to_field,       "Конец периода"),
            (self.btn_scan_folders,    "Открыть диалог выбора подпапок"),
            (self.btn_clear_folders,   "Очистить фильтры папок"),
            (self.save_stats_cb,       "Сохранять отчёт по итогам обработки"),
            (self.report_dir_field,    "Папка для сохранения отчётов"),
            (self.btn_reset_report,    "Сбросить папку отчётов"),
            (self.start_btn,           "Начать обработку (F5)"),
            (self.btn_no,              "Пропустить текущий файл (N / ←)"),
            (self.btn_yes,             "Сохранить текущий файл (Y / →)"),
            (self.btn_stop,            "Завершить и показать итоги (Esc)"),
            (self.btn_all,             "Сохранить все оставшиеся файлы (A)"),
        ]
        self._tooltip_targets = [(c, t) for c, t in items if c is not None]

    def _apply_tooltips(self):
        enabled = bool(self.settings.get("show_tooltips", True))
        for ctrl, txt in self._tooltip_targets:
            try:
                ctrl.tooltip = txt if enabled else None
                ctrl.update()
            except Exception:
                pass

    def _on_show_tooltips(self, value):
        self.settings["show_tooltips"] = bool(value)
        save_settings(self.settings)
        self._apply_tooltips()

    # ---------- НАСТРОЙКИ / СПРАВКА / О ПРОГРАММЕ ----------
    def _open_settings_dialog(self, e=None):
        theme_dd = _dd(
            value=self.settings.get("theme","system"),
            options=[
                ft.dropdown.Option(key="light", text="Светлая"),
                ft.dropdown.Option(key="dark", text="Тёмная"),
                ft.dropdown.Option(key="system", text="Системная"),
            ],
            on_change=self._on_theme_change,
            expand=True,
        )

        cb_tooltips = _cb(
            label="Показывать всплывающие подсказки",
            value=self.settings.get("show_tooltips", True),
            on_change=lambda ev: self._on_show_tooltips(_ev(ev)),
        )

        def ok(ev):
            self._close_dlg(dlg)

        content = ft.Container(
            content=ft.Column([
                ft.Text("Оформление", size=12, weight=ft.FontWeight.BOLD),
                ft.Container(content=theme_dd, padding=_pad_only(left=4, right=4)),
                self._spacer(),
                ft.Text("Поведение интерфейса", size=12, weight=ft.FontWeight.BOLD),
                ft.Container(content=cb_tooltips, padding=_pad_only(left=4)),
                ft.Container(content=ft.Text(
                    "Подсказки появляются при наведении мыши на элементы управления.",
                    size=HINT_SIZE, color=C_ON_VAR, italic=True),
                    padding=_pad_only(left=28, top=2)),
            ], spacing=GAP_FIELD, tight=True,
               horizontal_alignment=ft.CrossAxisAlignment.STRETCH),
            width=460, padding=4,
        )

        dlg = ft.AlertDialog(
            modal=True,
            title=ft.Row([
                ft.Icon(getattr(ft.Icons,"SETTINGS",None) or ft.Icons.SETTINGS,
                        color=ft.Colors.PRIMARY, size=20),
                ft.Text("Настройки", size=15, weight=ft.FontWeight.BOLD),
            ], spacing=8, tight=True),
            content=content,
            actions=[ft.FilledButton("OK", autofocus=True, on_click=ok)],
        )
        self._open_dlg(dlg)

    def _help_text(self):
        return (
            "КАК ПОЛЬЗОВАТЬСЯ\n"
            "════════════════════════════════════════════════════\n\n"
            "1. ВЫБЕРИТЕ ПАПКУ\n"
            "   • Нажмите «Выбрать папку» и укажите каталог с файлами, либо\n"
            "     вставьте путь в поле вручную и нажмите Enter.\n"
            "   • Подпапки сканируются автоматически.\n\n"
            "2. НАСТРОЙТЕ ФОРМАТЫ\n"
            "   • Выберите набор входных расширений (Только PSD, растровые,\n"
            "     всё поддерживаемое) или нажмите «Настроить» и отметьте вручную.\n"
            "   • Отметьте форматы вывода: PNG (прозрачный/белый), JPG, WEBP,\n"
            "     TIFF, BMP, PDF, AVIF. Можно выбрать несколько одновременно.\n"
            "   • Определите поведение, если файл уже существует:\n"
            "     перезаписывать или пропускать.\n\n"
            "3. ФИЛЬТРЫ (необязательно)\n"
            "   • Поиск по артикулам:\n"
            "     – «Совпадение»: полное (имя файла = артикул) или частичное.\n"
            "     – Артикулы можно ввести списком, загрузить из TXT/CSV/XLSX/DOCX\n"
            "       или из Google Sheets.\n"
            "   • Фильтр папок: маски через запятую, исключения для ненужных.\n"
            "   • Фильтр по дате: за сутки / неделю / месяц / год / свой период.\n"
            "   • Изменение фильтров обновляет превью сразу, но не запускает\n"
            "     и не завершает обработку. Для обработки нажмите СТАРТ.\n\n"
            "4. СТАРТ\n"
            "   • Нажмите «СТАРТ» (F5) — начнётся обработка по очереди.\n"
            "   • Управление: Y / → — сохранить, N / ← — пропустить,\n"
            "     A — сохранить все оставшиеся, Esc — завершить и показать итог.\n"
            "   • После последнего файла отчёт НЕ появляется автоматически.\n"
            "     Нажмите «Стоп / Итог» (Esc), чтобы увидеть итоги.\n\n"
            "5. ОТЧЁТЫ\n"
            "   • Галочка «Сохранять отчёт» создаёт TXT-отчёт при нажатии «Стоп».\n\n"
            "ОБЩИЕ НАСТРОЙКИ\n"
            "────────────────────────────────────────────────────\n"
            "   Кнопка «Настройки» открывает диалог с выбором темы\n"
            "   оформления и переключателем всплывающих подсказок.\n\n"
            "СОКРАЩЕНИЯ\n"
            "────────────────────────────────────────────────────\n"
            "   F5 — старт\n"
            "   Y / → — сохранить\n"
            "   N / ← — пропустить\n"
            "   A — сохранить все\n"
            "   Esc — завершить\n"
        )

    def _open_help_dialog(self, e=None):
        tf = ft.Text(self._help_text(), size=12,
                      font_family="Consolas", selectable=True)
        body = ft.Container(content=ft.Column([tf], scroll=ft.ScrollMode.AUTO,
                                                tight=True),
                             height=480, width=640, padding=6)
        dlg = ft.AlertDialog(
            modal=True,
            title=ft.Row([
                ft.Icon(getattr(ft.Icons,"HELP_OUTLINE",None) or ft.Icons.HELP,
                        color=ft.Colors.PRIMARY, size=20),
                ft.Text("Справка", size=15, weight=ft.FontWeight.BOLD),
            ], spacing=8, tight=True),
            content=body,
            actions=[ft.FilledButton("OK",
                     on_click=lambda ev: self._close_dlg(dlg))],
        )
        self._open_dlg(dlg)

    def _open_about_dialog(self, e=None):
        def li(txt):
            return ft.Row([
                ft.Text("•", color=ft.Colors.PRIMARY, weight=ft.FontWeight.BOLD),
                ft.Text(txt, size=12, selectable=False),
            ], spacing=8, tight=True)

        header = ft.Column([
            ft.Text(APP_NAME, size=22, weight=ft.FontWeight.BOLD,
                     color=ft.Colors.PRIMARY),
            ft.Text(f"Версия {APP_VERSION}", size=11, color=C_ON_VAR),
            ft.Text(APP_TAGLINE, size=12, italic=True),
        ], spacing=2, tight=True)

        features = ft.Column([
            ft.Text("Возможности:", size=12, weight=ft.FontWeight.BOLD),
            li("Поддержка PSD, PSB, TIFF, TGA, BMP, WebP, GIF, ICO, PNG, JPEG"),
            li("Импорт PDF и Adobe Illustrator (первая страница)"),
            li("HEIC / HEIF через pillow-heif"),
            li("Экспорт: PNG (прозр./бел.), JPG, WEBP, TIFF, BMP, PDF, AVIF"),
            li("Пакетная обработка с пошаговым превью"),
            li("Фильтры по артикулам, папкам и датам"),
            li("Импорт артикулов из TXT / CSV / XLSX / DOCX и Google Sheets"),
            li("Отчёты по завершении работы"),
            li("Тёмная и светлая тема, адаптивная раскладка"),
        ], spacing=3, tight=True)

        footer_items = [
            ft.Divider(height=1),
            ft.Row([
                ft.Icon(getattr(ft.Icons,"LINK",None) or ft.Icons.LINK,
                        size=14, color=C_ON_VAR),
                ft.Text(APP_REPO, size=11, color=ft.Colors.PRIMARY,
                         selectable=True),
            ], spacing=6, tight=True),
            ft.Text(f"Лицензия: {APP_LICENSE}",
                     size=11, italic=True, color=C_ON_VAR),
            ft.Text(f"© {APP_YEAR} {APP_AUTHOR}",
                     size=11, italic=True, color=C_ON_VAR),
        ]
        footer = ft.Column(footer_items, spacing=6, tight=True)

        content = ft.Container(
            content=ft.Column([header, ft.Divider(height=1), features, footer],
                                spacing=GAP, tight=True),
            width=520, padding=4,
        )

        dlg = ft.AlertDialog(
            modal=True,
            title=ft.Row([
                ft.Icon(getattr(ft.Icons,"INFO_OUTLINE",None) or ft.Icons.INFO,
                        color=ft.Colors.PRIMARY, size=20),
                ft.Text("О программе", size=15, weight=ft.FontWeight.BOLD),
            ], spacing=8, tight=True),
            content=content,
            actions=[
                ft.FilledButton("OK", autofocus=True,
                                 on_click=lambda ev: self._close_dlg(dlg)),
            ],
        )
        self._open_dlg(dlg)

    def _build_ui(self):
        self._left_scroll_indicator = ft.Container(
            width=4, bgcolor=ft.Colors.PRIMARY, border_radius=2,
            height=0, visible=False,
        )
        self._left_scroll_grip = ft.GestureDetector(
            content=ft.Container(
                content=self._left_scroll_indicator,
                padding=_pad_sym(3, 0),
                bgcolor="transparent",
            ),
            mouse_cursor=ft.MouseCursor.CLICK,
            on_pan_start=self._on_scroll_drag_start,
            on_pan_update=self._on_scroll_drag,
        )
        self._left_scroll_indicator_holder = ft.Container(
            content=self._left_scroll_grip,
            width=10, bgcolor="transparent",
            alignment=TOP_CENTER,
        )

        left_col = _safe(
            ft.Column,
            spacing=GAP,
            scroll=ft.ScrollMode.HIDDEN,
            expand=True,
            on_scroll=self._on_left_scroll,
            horizontal_alignment=ft.CrossAxisAlignment.STRETCH,
        )
        self._left_col = left_col

        left_inner = ft.Container(
            content=ft.Row(
                [ft.Container(content=left_col, expand=True, padding=_pad_only(right=4)),
                 self._left_scroll_indicator_holder],
                spacing=0, expand=True,
                vertical_alignment=ft.CrossAxisAlignment.STRETCH,
            ),
            padding=0,
            expand=True,
        )

        self._left_box = ft.Container(
            content=left_inner,
            width=self._left_panel_width,
            padding=_pad_only(left=2, top=2, bottom=2, right=2),
        )

        # ---- Кнопки верхней панели "Общие настройки" ----
        self.btn_settings = ft.OutlinedButton(
            "Настройки",
            icon=getattr(ft.Icons,"SETTINGS",None) or ft.Icons.SETTINGS,
            height=34, on_click=self._open_settings_dialog,
            style=_btn_style())
        self.btn_help = ft.OutlinedButton(
            "Справка", icon=getattr(ft.Icons,"HELP_OUTLINE",ft.Icons.HELP),
            height=34, on_click=self._open_help_dialog,
            style=_btn_style())
        self.btn_about = ft.OutlinedButton(
            "О программе", icon=getattr(ft.Icons,"INFO_OUTLINE",ft.Icons.INFO),
            height=34, on_click=self._open_about_dialog,
            style=_btn_style())

        self._top_row = ft.Row(
            [self.btn_settings, self.btn_help, self.btn_about],
            spacing=GAP_FIELD, wrap=True, run_spacing=GAP_FIELD,
            alignment=ft.MainAxisAlignment.START,
            run_alignment=ft.MainAxisAlignment.START,
            vertical_alignment=ft.CrossAxisAlignment.CENTER,
        )

        self._top_card = ft.Container(
            content=ft.Column([
                ft.Text("Общие настройки программы", size=12, weight=ft.FontWeight.BOLD),
                self._top_row,
            ], spacing=GAP_FIELD, tight=True,
               horizontal_alignment=ft.CrossAxisAlignment.STRETCH),
            padding=8,
            bgcolor=C_SURF_HIGH, border_radius=14,
        )

        # ---- Карточка "Папка с исходниками" ----
        self.folder_path = _tf(value=self.settings.get("last_folder") or "",
                                on_submit=self._commit_typed_path,
                                on_blur=self._commit_typed_path, expand=True)
        if self.settings.get("last_folder"):
            p = Path(self.settings["last_folder"])
            if p.is_dir():
                self.current_folder = p

        self.btn_choose_folder = ft.FilledButton(
            "Выбрать папку", icon=ft.Icons.FOLDER_OPEN,
            height=ROW_H, on_click=self._choose_folder,
            style=_btn_style())
        self.btn_rescan = ft.OutlinedButton(
            "Пересканировать", icon=ft.Icons.REFRESH,
            height=ROW_H, on_click=self._rescan,
            style=_btn_style())
        self.btn_cancel_scan = ft.OutlinedButton(
            "Прервать скан", icon=ft.Icons.STOP,
            height=ROW_H, on_click=self._cancel_scan,
            style=_btn_style(), disabled=True)

        self._folder_title = ft.Text("Папка с исходниками", size=13, weight=ft.FontWeight.BOLD,
                                       no_wrap=True, overflow=ft.TextOverflow.ELLIPSIS)

        self._folder_pair = ft.Container(
            content=ft.Row(
                [self.btn_choose_folder, self.folder_path],
                spacing=GAP,
                vertical_alignment=ft.CrossAxisAlignment.CENTER,
            ),
        )

        folder_card = self._card(
            ft.Row([self._folder_title], spacing=GAP,
                   vertical_alignment=ft.CrossAxisAlignment.CENTER),
            self._folder_pair,
            ft.Row([self.btn_rescan, self.btn_cancel_scan],
                   spacing=GAP, wrap=True, run_spacing=GAP,
                   run_alignment=ft.MainAxisAlignment.START,
                   vertical_alignment=ft.CrossAxisAlignment.CENTER),
        )

        # ---- Карточка "Форматы и поведение" ----
        self.preset_dd = _dd(value=self._current_preset(),
                              options=[ft.dropdown.Option(k) for k in INPUT_PRESETS]+
                                      [ft.dropdown.Option(INPUT_PRESET_CUSTOM)],
                              on_change=self._on_preset_change, expand=True)
        self.btn_setup = ft.OutlinedButton("Настроить", height=ROW_H,
                                            on_click=self._open_exts_dialog,
                                            style=_btn_style())
        self.exts_summary = ft.Text(self._exts_summary(), size=HINT_SIZE, color=C_ON_VAR,
                                      no_wrap=True, overflow=ft.TextOverflow.ELLIPSIS)

        out_boxes = []
        self._out_containers = []
        for fid,label,_ in OUTPUT_FORMATS:
            cb = _cb(value=self.out_vars[fid],
                      on_change=lambda e,fid=fid: self._on_output_toggle(fid,_ev(e)))
            lbl = ft.Text(label, size=TXT_SIZE, color=C_ON_SURF,
                           no_wrap=False, max_lines=2, expand=True)
            cont = ft.Container(
                content=ft.Row([cb, lbl], spacing=4, tight=True,
                                vertical_alignment=ft.CrossAxisAlignment.CENTER),
                col={"xs": 12, "sm": 6, "md": 4, "lg": 3, "xl": 3})
            self._out_containers.append(cont)
            out_boxes.append(cont)
        self.out_grid = ft.ResponsiveRow(
            out_boxes, spacing=GAP, run_spacing=GAP_FIELD)

        self.overwrite_dd = _dd(
            value="Перезаписывать" if self.settings.get("overwrite_mode") == "overwrite" else "Пропускать",
            options=[ft.dropdown.Option("Перезаписывать"),ft.dropdown.Option("Пропускать")],
            on_change=self._on_overwrite_change, expand=True)

        formats_body = [
            ft.Row([self.preset_dd, self.btn_setup],
                   spacing=GAP, vertical_alignment=ft.CrossAxisAlignment.CENTER),
            self.exts_summary,
            self._spacer(),
            ft.Text("Форматы вывода:", size=LBL_SIZE, weight=LBL_WEIGHT, color=C_ON_SURF),
            self.out_grid,
            self._spacer(),
            self._field("Если файл уже есть:", self.overwrite_dd),
        ]
        self.sec_formats = Section("Форматы и поведение", formats_body,
                                    "sec_open_formats",self,
                                    initially_open=self.settings.get("sec_open_formats",True))

        # ---- Карточка "Поиск по артикулам" ----
        self.articles_text = _tf_multi(value=self.settings.get("articles",""),
                                        hint_text="По одному в строке или через запятую",
                                        on_change=self._on_articles_change)
        self.match_dd = _dd(value="Частичное" if self.settings.get("match_mode") == "partial" else "Полное",
                             options=[ft.dropdown.Option("Частичное"),ft.dropdown.Option("Полное")],
                             on_change=self._on_match_change, expand=True)

        self.btn_from_file = ft.OutlinedButton(
            "Из файла…", icon=ft.Icons.UPLOAD_FILE,
            height=ROW_H, on_click=self._load_articles_from_file,
            style=_btn_style())
        self.btn_from_gsheet = ft.OutlinedButton(
            "Google Sheets…", icon=ft.Icons.CLOUD_DOWNLOAD,
            height=ROW_H, on_click=self._load_articles_from_gsheet,
            style=_btn_style())
        self.btn_clear_search = ft.OutlinedButton(
            "Очистить", height=ROW_H, on_click=self._clear_search,
            style=_btn_style())
        self.btn_apply_filter = ft.FilledButton(
            "Применить", height=ROW_H, on_click=self._apply_filter_now,
            style=_btn_style())

        articles_body = [
            self._field("Совпадение:", self.match_dd),
            self._spacer(),
            self._field("Артикулы (пусто — без фильтра):", self.articles_text),
            self._spacer(),
            ft.Row([
                self.btn_from_file, self.btn_from_gsheet,
                self.btn_clear_search, self.btn_apply_filter,
            ], spacing=GAP, wrap=True, run_spacing=GAP,
               run_alignment=ft.MainAxisAlignment.START,
               vertical_alignment=ft.CrossAxisAlignment.CENTER),
        ]
        self.sec_articles = Section("Поиск по артикулам", articles_body,
                                     "sec_open_articles",self,
                                     initially_open=self.settings.get("sec_open_articles",True))

        # ---- Карточка "Фильтр папок и даты" ----
        self.include_text = _tf_multi(
            value=self.settings.get("include_folders",""),
            hint_text="Render, PSD*  —  через запятую или в столбик",
            on_change=self._on_filters_change)
        self.exclude_text = _tf_multi(
            value=self.settings.get("exclude_folders",""),
            hint_text="old, backup  —  через запятую или в столбик",
            on_change=self._on_filters_change)
        date_names = {"all":"Все","day":"За сутки","week":"За неделю",
                      "month":"За месяц","year":"За год","custom":"Свой период"}
        self.date_dd = _dd(value=date_names.get(self.settings.get("date_mode","all"),"Все"),
                            options=[ft.dropdown.Option(x) for x in date_names.values()],
                            on_change=self._on_date_mode_change, expand=True)
        self.date_kind_dd = _dd(value="Изменения файла" if self.settings.get("date_kind") == "modified" else "Создания файла",
                                 options=[ft.dropdown.Option("Изменения файла"),ft.dropdown.Option("Создания файла")],
                                 on_change=self._on_date_kind_change, expand=True)
        self.date_from_field = _tf(value=self.settings.get("date_from",""),
                                    hint_text="ДД.ММ.ГГГГ",
                                    on_change=self._on_filters_change)
        self.date_to_field = _tf(value=self.settings.get("date_to",""),
                                  hint_text="ДД.ММ.ГГГГ",
                                  on_change=self._on_filters_change)

        self.btn_scan_folders = ft.FilledButton(
            "Отметить подпапки…", height=ROW_H,
            on_click=self._scan_folders_dialog,
            style=_btn_style())
        self.btn_clear_folders = ft.OutlinedButton(
            "Очистить", height=ROW_H,
            on_click=self._clear_folder_filter,
            style=_btn_style())

        filters_body = [
            self._field("Искать только в папках:", self.include_text,
                        hint="Пусто — поиск во всех"),
            self._spacer(),
            self._field("Не искать в папках:", self.exclude_text,
                        hint="Маски: Render, *_final, *backup*"),
            self._spacer(),
            ft.Row([self.btn_scan_folders, self.btn_clear_folders],
                   spacing=GAP, wrap=True, run_spacing=GAP,
                   run_alignment=ft.MainAxisAlignment.START,
                   vertical_alignment=ft.CrossAxisAlignment.CENTER),
            self._spacer(),
            ft.Divider(height=1),
            ft.Text("Дата", size=LBL_SIZE, weight=LBL_WEIGHT, color=C_ON_SURF),
            self._field("Период:", self.date_dd),
            self._field("Считать:", self.date_kind_dd),
            self._field(
                "Свой период:",
                ft.ResponsiveRow([
                    ft.Container(
                        content=ft.Column([
                            ft.Text("с", size=HINT_SIZE, color=C_ON_VAR),
                            self.date_from_field,
                        ], spacing=2, tight=True),
                        col={"xs": 12, "sm": 6, "md": 6, "lg": 6},
                    ),
                    ft.Container(
                        content=ft.Column([
                            ft.Text("по", size=HINT_SIZE, color=C_ON_VAR),
                            self.date_to_field,
                        ], spacing=2, tight=True),
                        col={"xs": 12, "sm": 6, "md": 6, "lg": 6},
                    ),
                ], spacing=GAP, run_spacing=GAP_FIELD)
            ),
        ]
        self.sec_filters = Section("Фильтр папок и даты", filters_body,
                                    "sec_open_filters",self,
                                    initially_open=self.settings.get("sec_open_filters",True))

        # ---- Карточка "Статистика и отчёты" ----
        self.save_stats_cb = _cb(label="Сохранять отчёт",
                                  value=self.settings.get("save_stats",False),
                                  on_change=lambda e: self._on_save_stats(_ev(e)))
        self.report_dir_field = _tf(value=self.settings.get("report_dir",""),
                                     on_change=self._on_report_dir_change, expand=True)
        self.report_hint = ft.Text(self._report_hint_text(), size=HINT_SIZE, color=C_ON_VAR,
                                     italic=True, no_wrap=False)
        self.btn_reset_report = ft.OutlinedButton(
            "По умолчанию", height=ROW_H,
            on_click=self._reset_report_dir,
            style=_btn_style())
        stats_body = [
            self.save_stats_cb,
            self._spacer(),
            self._field("Папка для отчётов:",
                        ft.Row([self.report_dir_field,
                                ft.IconButton(icon=ft.Icons.FOLDER_OPEN, icon_size=20,
                                               on_click=self._choose_report_dir,
                                               tooltip="Выбрать папку")],
                               spacing=GAP, vertical_alignment=ft.CrossAxisAlignment.CENTER)),
            ft.Row([self.btn_reset_report], spacing=GAP, wrap=True),
            self.report_hint,
        ]
        self.sec_stats = Section("Статистика и отчёты", stats_body,
                                  "sec_open_stats",self,
                                  initially_open=self.settings.get("sec_open_stats",True))

        left_col.controls.extend([
            self._top_card, folder_card, self.sec_formats, self.sec_articles,
            self.sec_filters, self.sec_stats,
        ])

        self.sash_inner = ft.Container(
            width=3, bgcolor=ft.Colors.OUTLINE_VARIANT, border_radius=2, expand=True,
        )
        self.sash = ft.GestureDetector(
            content=ft.Container(
                content=self.sash_inner,
                padding=_pad_sym(5, 0),
                bgcolor="transparent",
                expand=True,
            ),
            mouse_cursor=ft.MouseCursor.RESIZE_LEFT_RIGHT,
            on_pan_start=self._on_sash_start,
            on_pan_update=self._on_sash_drag,
            on_pan_end=self._on_sash_end,
            on_hover=self._on_sash_hover,
        )

        self.dark_switch = ThemeToggle(
            value=self.settings.get("dark_preview",False),
            on_change=self._on_preview_bg_change,
        )
        self.counter_top = ft.Text("", size=TXT_SIZE, color=ft.Colors.PRIMARY,
                                     weight=ft.FontWeight.BOLD)

        self.start_btn = ft.FilledButton(
            "СТАРТ   (F5)", icon=ft.Icons.PLAY_ARROW, height=ROW_H_MAIN,
            on_click=self._start_work,
            style=_btn_style(bgcolor=ft.Colors.GREEN_600, color=ft.Colors.WHITE,
                              text_style=ft.TextStyle(size=16,weight=ft.FontWeight.BOLD)))
        self.start_box = ft.Container(content=ft.Row([self.start_btn],
                                                       alignment=ft.MainAxisAlignment.CENTER))

        self.preview_img = ft.Image(src="", fit=BOX_CONTAIN, expand=True, visible=False)
        self.preview_hint = ft.Text("Превью появится здесь", size=12, color=C_ON_VAR)
        self.preview_hint_box = ft.Container(content=self.preview_hint,
                                              alignment=CENTER, expand=True)
        self.preview_container = ft.Container(
            content=ft.Stack([
                self.preview_hint_box,
                ft.Container(content=self.preview_img, alignment=CENTER, expand=True),
            ], expand=True),
            expand=True, bgcolor=C_SURF_HIGH, border_radius=14, padding=PAD,
        )
        self.preview_bg = C_SURF_HIGH

        self.btn_no = ft.FilledButton("Нет   (N / ←)", height=54,
            on_click=self._reject, disabled=True,
            style=_btn_style(bgcolor=ft.Colors.RED, color=ft.Colors.WHITE,
                              text_style=ft.TextStyle(size=14,weight=ft.FontWeight.BOLD)))
        self.btn_yes = ft.FilledButton("Да   (Y / →)", height=54,
            on_click=self._accept, disabled=True,
            style=_btn_style(bgcolor=ft.Colors.GREEN, color=ft.Colors.WHITE,
                              text_style=ft.TextStyle(size=14,weight=ft.FontWeight.BOLD)))
        self.btn_stop = ft.OutlinedButton("Стоп / Итог   (Esc)", height=ROW_H_MAIN,
            on_click=self._finish, disabled=True,
            style=_btn_style(text_style=ft.TextStyle(size=TXT_SIZE)))
        self.btn_all = ft.OutlinedButton("Да для всех   (A)", height=ROW_H_MAIN,
            on_click=self._accept_all, disabled=True,
            style=_btn_style(text_style=ft.TextStyle(size=TXT_SIZE)))

        self.status_text = ft.Text("", size=TXT_SIZE, color=C_ON_VAR, text_align=ft.TextAlign.CENTER)
        self.name_text = ft.Text("", size=12, weight=ft.FontWeight.BOLD,
                                  text_align=ft.TextAlign.CENTER,
                                  max_lines=2, overflow=ft.TextOverflow.ELLIPSIS)

        sep = ft.Text("·", size=HINT_SIZE, color=C_ON_VAR)
        self.hot_bar = ft.Container(
            content=ft.Row([
                self._hk("N / ←", "Нет"),
                sep,
                self._hk("Y / →", "Да"),
                sep,
                self._hk("A", "Для всех"),
                sep,
                self._hk("Esc", "Стоп"),
                sep,
                self._hk("F5", "Старт"),
            ], spacing=6, wrap=True, run_spacing=4,
               alignment=ft.MainAxisAlignment.CENTER,
               run_alignment=ft.MainAxisAlignment.CENTER),
            padding=6, bgcolor=C_SURF_HIGHEST, border_radius=8,
        )

        self.preview_toggle_row = ft.Row([
            ft.Text("Светлый", size=TXT_SIZE, color=C_ON_VAR),
            self.dark_switch,
            ft.Text("Тёмный", size=TXT_SIZE, color=C_ON_VAR),
            ft.Container(expand=True),
            self.counter_top,
        ], spacing=GAP, vertical_alignment=ft.CrossAxisAlignment.CENTER)

        self.btn_row_main = ft.Row([self.btn_no, self.btn_yes],
                                    spacing=GAP,
                                    wrap=True, run_spacing=GAP,
                                    alignment=ft.MainAxisAlignment.CENTER,
                                    run_alignment=ft.MainAxisAlignment.CENTER,
                                    vertical_alignment=ft.CrossAxisAlignment.CENTER)
        self.btn_row_extra = ft.Row([self.btn_stop, self.btn_all],
                                     spacing=GAP,
                                     wrap=True, run_spacing=GAP,
                                     alignment=ft.MainAxisAlignment.CENTER,
                                     run_alignment=ft.MainAxisAlignment.CENTER,
                                     vertical_alignment=ft.CrossAxisAlignment.CENTER)

        self.right_col = ft.Column([
            self.preview_toggle_row,
            self.start_box,
            self.preview_container,
            self.btn_row_main,
            self.btn_row_extra,
            ft.Row([self.status_text], alignment=ft.MainAxisAlignment.CENTER),
            ft.Row([self.name_text], alignment=ft.MainAxisAlignment.CENTER),
            self.hot_bar,
        ], spacing=GAP, expand=True,
           horizontal_alignment=ft.CrossAxisAlignment.STRETCH)
        self._right_box = ft.Container(content=self.right_col, expand=True,
                                        padding=_pad_only(left=2))

        layout = ft.Row(
            [self._left_box, self.sash, self._right_box],
            expand=True, spacing=0,
            vertical_alignment=ft.CrossAxisAlignment.STRETCH,
        )
        self.page.add(layout)
        self._apply_preview_bg()

    # ------- АДАПТИВНОСТЬ ЛЕВОЙ ПАНЕЛИ -------
    def _set_out_cols(self, cols):
        for c in self._out_containers:
            try:
                c.col = cols
                c.update()
            except Exception:
                pass

    def _apply_adaptive(self, left_w):
        narrow       = left_w < NARROW_W
        very_narrow  = left_w < VERY_NARROW_W
        ultra_narrow = left_w < 290

        try:
            self._set_folder_pair_stacked(very_narrow)
        except Exception:
            pass

        try:
            if narrow:
                self._set_top_row_alignment(ft.MainAxisAlignment.START)
            else:
                self._set_top_row_alignment(ft.MainAxisAlignment.END)
        except Exception:
            pass

        try:
            if ultra_narrow:
                self.btn_choose_folder.text = ""
                self.btn_choose_folder.width = 44
            else:
                self.btn_choose_folder.text = "Выбрать папку"
                self.btn_choose_folder.width = None
            self.btn_choose_folder.update()
        except Exception:
            pass

        try:
            if ultra_narrow:
                self.btn_rescan.text = ""
                self.btn_cancel_scan.text = ""
            elif narrow:
                self.btn_rescan.text = "Скан"
                self.btn_cancel_scan.text = "Стоп"
            else:
                self.btn_rescan.text = "Пересканировать"
                self.btn_cancel_scan.text = "Прервать скан"
            self.btn_rescan.update()
            self.btn_cancel_scan.update()
        except Exception:
            pass

        try:
            if very_narrow:
                self.btn_setup.text = "…"
            elif narrow:
                self.btn_setup.text = "Настр."
            else:
                self.btn_setup.text = "Настроить"
            self.btn_setup.update()
        except Exception:
            pass

        try:
            if ultra_narrow:
                self.btn_settings.text = ""
                self.btn_help.text = ""
                self.btn_about.text = ""
            elif very_narrow:
                self.btn_settings.text = "⚙"
                self.btn_help.text = "?"
                self.btn_about.text = "i"
            elif narrow:
                self.btn_settings.text = "Настр."
                self.btn_help.text = "Справка"
                self.btn_about.text = "О прогр."
            else:
                self.btn_settings.text = "Настройки"
                self.btn_help.text = "Справка"
                self.btn_about.text = "О программе"
            self.btn_settings.update()
            self.btn_help.update()
            self.btn_about.update()
        except Exception:
            pass

        try:
            self.exts_summary.visible = not narrow
            self.exts_summary.update()
        except Exception:
            pass

        try:
            if very_narrow:
                self._set_out_cols({"xs": 12, "sm": 12, "md": 12, "lg": 12, "xl": 12})
            elif narrow:
                self._set_out_cols({"xs": 6, "sm": 6, "md": 6, "lg": 6, "xl": 6})
            else:
                self._set_out_cols({"xs": 12, "sm": 6, "md": 4, "lg": 3, "xl": 3})
        except Exception:
            pass

        try:
            if ultra_narrow:
                self.btn_from_file.text = ""
                self.btn_from_gsheet.text = ""
            elif very_narrow:
                self.btn_from_file.text = "Файл"
                self.btn_from_gsheet.text = "Sheets"
            elif narrow:
                self.btn_from_file.text = "Файл…"
                self.btn_from_gsheet.text = "Sheets…"
            else:
                self.btn_from_file.text = "Из файла…"
                self.btn_from_gsheet.text = "Google Sheets…"
            self.btn_from_file.update()
            self.btn_from_gsheet.update()
        except Exception:
            pass

        try:
            if ultra_narrow:
                self.btn_clear_search.text = "✕"
                self.btn_apply_filter.text = "ОК"
            elif narrow:
                self.btn_clear_search.text = "Сброс"
                self.btn_apply_filter.text = "ОК"
            else:
                self.btn_clear_search.text = "Очистить"
                self.btn_apply_filter.text = "Применить"
            self.btn_clear_search.update()
            self.btn_apply_filter.update()
        except Exception:
            pass

        try:
            if ultra_narrow:
                self.btn_scan_folders.text = "Папки"
                self.btn_clear_folders.text = "✕"
            elif very_narrow:
                self.btn_scan_folders.text = "Подпапки"
                self.btn_clear_folders.text = "Сброс"
            elif narrow:
                self.btn_scan_folders.text = "Подпапки…"
                self.btn_clear_folders.text = "Сброс"
            else:
                self.btn_scan_folders.text = "Отметить подпапки…"
                self.btn_clear_folders.text = "Очистить"
            self.btn_scan_folders.update()
            self.btn_clear_folders.update()
        except Exception:
            pass

        try:
            if very_narrow:
                self.btn_reset_report.text = "↺"
            elif narrow:
                self.btn_reset_report.text = "Сброс"
            else:
                self.btn_reset_report.text = "По умолчанию"
            self.btn_reset_report.update()
        except Exception:
            pass

        try:
            self.report_hint.visible = not narrow
            self.report_hint.update()
        except Exception:
            pass

        try:
            for sec in (self.sec_formats, self.sec_articles,
                        self.sec_filters, self.sec_stats):
                sec.set_compact(narrow)
        except Exception:
            pass

        try:
            new_pad = 8 if very_narrow else (10 if narrow else PAD)
            for sec in (self.sec_formats, self.sec_articles,
                        self.sec_filters, self.sec_stats):
                sec.set_body_padding(new_pad)
        except Exception:
            pass

        try:
            if narrow:
                self.sec_formats.set_title("Форматы")
                self.sec_articles.set_title("Артикулы")
                self.sec_filters.set_title("Фильтр папок")
                self.sec_stats.set_title("Отчёты")
            else:
                self.sec_formats.set_title("Форматы и поведение")
                self.sec_articles.set_title("Поиск по артикулам")
                self.sec_filters.set_title("Фильтр папок и даты")
                self.sec_stats.set_title("Статистика и отчёты")
        except Exception:
            pass

        try:
            self.page.update()
        except Exception:
            pass

    def _apply_right_adaptive(self, right_w):
        try:
            self.hot_bar.visible = right_w >= RIGHT_NARROW
            self.hot_bar.update()
        except Exception:
            pass

        try:
            if right_w < 340:
                self.btn_no.text = "Нет"
                self.btn_yes.text = "Да"
                self.btn_stop.text = "Стоп"
                self.btn_all.text = "Все"
            elif right_w < 480:
                self.btn_no.text = "Нет"
                self.btn_yes.text = "Да"
                self.btn_stop.text = "Стоп / Итог"
                self.btn_all.text = "Для всех"
            else:
                self.btn_no.text = "Нет   (N / ←)"
                self.btn_yes.text = "Да   (Y / →)"
                self.btn_stop.text = "Стоп / Итог   (Esc)"
                self.btn_all.text = "Да для всех   (A)"
            self.btn_no.update()
            self.btn_yes.update()
            self.btn_stop.update()
            self.btn_all.update()
        except Exception:
            pass

    def _on_left_scroll(self, e):
        try:
            pos = float(getattr(e, "pixels", 0) or 0)
            mx = float(getattr(e, "max_scroll_extent", 0) or 0)
            self._left_scroll_pos = pos
            self._left_scroll_max = mx
            self._update_left_scroll_indicator()
        except Exception:
            pass

    def _on_scroll_drag_start(self, e):
        try:
            self._scroll_drag_start = float(getattr(e, "global_y", 0) or 0)
        except Exception:
            self._scroll_drag_start = None

    def _on_scroll_drag(self, e):
        try:
            gy = getattr(e, "global_y", None)
            if gy is None or self._scroll_drag_start is None:
                return
            if self._left_scroll_max <= 0:
                return
            holder_h = self._left_scroll_indicator_holder.height or 600
            delta_px = float(gy) - self._scroll_drag_start
            self._scroll_drag_start = float(gy)
            ratio = self._left_scroll_max / max(1.0, holder_h)
            new_pos = self._left_scroll_pos + delta_px * ratio
            new_pos = max(0.0, min(self._left_scroll_max, new_pos))
            self._left_scroll_pos = new_pos
            try:
                self._left_col.scroll_to(offset=int(new_pos))
            except Exception:
                pass
            self._update_left_scroll_indicator()
        except Exception:
            pass

    def _update_left_scroll_indicator(self):
        try:
            holder_h = self._left_scroll_indicator_holder.height or 0
            if holder_h <= 0:
                holder_h = (self.page.window.height or 900) - 40
            visible_h = max(40, int(holder_h * 0.25))
            if self._left_scroll_max > 0:
                travel = max(0, holder_h - visible_h)
                ratio = min(1.0, max(0.0, self._left_scroll_pos / self._left_scroll_max))
                top = int(travel * ratio)
            else:
                top = 0
            self._left_scroll_indicator.height = visible_h
            self._left_scroll_indicator.visible = self._left_scroll_max > 0
            self._left_scroll_indicator_holder.alignment = TOP_CENTER
            self._left_scroll_indicator_holder.content = ft.Container(
                content=self._left_scroll_grip,
                padding=_pad_only(top=top),
                width=10, bgcolor="transparent",
            )
            self._left_scroll_indicator_holder.update()
        except Exception:
            pass

    def _get_global_x(self, e):
        for attr in ("global_x", "local_x"):
            v = getattr(e, attr, None)
            if v is not None:
                try:
                    return float(v)
                except Exception:
                    pass
        for attr in ("global_position", "local_position", "position"):
            p = getattr(e, attr, None)
            if p is not None:
                x = getattr(p, "x", None)
                if x is not None:
                    try:
                        return float(x)
                    except Exception:
                        pass
        return None

    def _on_sash_start(self, e):
        self._sash_start_w = self._left_panel_width
        self._sash_mouse_start_x = self._get_global_x(e)
        self._sash_last_gx = self._sash_mouse_start_x

    def _on_sash_end(self, e):
        self._sash_mouse_start_x = None
        self._sash_last_gx = None

    def _on_sash_drag(self, e):
        try:
            gx = self._get_global_x(e)
            dx = 0
            if gx is not None:
                if self._sash_mouse_start_x is not None:
                    dx = int(gx - self._sash_mouse_start_x)
                    self._sash_mouse_start_x = gx
            else:
                for attr in ("delta_x", "local_delta_x", "primary_delta_x"):
                    v = getattr(e, attr, None)
                    if v:
                        try:
                            dx = int(v)
                        except Exception:
                            dx = 0
                        break
            if dx == 0:
                return
            win_w = self._read_page_width_safe()
            max_left = win_w - SASH_W - RIGHT_MIN
            new_w = max(LEFT_MIN, min(max_left, self._left_panel_width + dx))
            if new_w != self._left_panel_width:
                self._left_panel_width = new_w
                self._left_box.width = new_w
                self.settings["left_panel_width"] = new_w
                try:
                    self._left_box.update()
                except Exception:
                    pass
                self._apply_adaptive(new_w)
                self._apply_right_adaptive(win_w - new_w - SASH_W)
        except Exception:
            pass

    def _on_sash_hover(self, e):
        try:
            is_hover = str(getattr(e,"data","")).lower() == "true"
            self.sash_inner.bgcolor = (ft.Colors.PRIMARY if is_hover
                                        else ft.Colors.OUTLINE_VARIANT)
            self.sash_inner.update()
        except Exception:
            pass

    def _on_resize(self, e):
        w = self._extract_width_from_event(e)
        if w:
            self._resize_event_width = w
            self._apply_width(w, force=True)
        else:
            self._apply_width(self._read_page_width_safe(), force=True)
        try:
            if self._resize_timer is not None:
                self._resize_timer.cancel()
        except Exception:
            pass
        self._resize_timer = threading.Timer(0.15, self._process_resize_delayed)
        self._resize_timer.daemon = True
        self._resize_timer.start()

    def _process_resize_delayed(self):
        try:
            self._resize_timer = None
        except Exception:
            pass
        try:
            w = self._resize_event_width
            if not w:
                w = self._read_page_width_safe()
            self._last_applied_width = None
            self._apply_width(w, force=True)
        except Exception:
            pass

    def _exts_summary(self):
        labels, seen = [], set()
        for ext in sorted(self._input_exts):
            g = INPUT_EXT_INFO.get(ext)
            if g and g not in seen:
                labels.append(g)
                seen.add(g)
            elif not g:
                labels.append(ext)
        n = len(self._input_exts)
        return f"Выбрано: {n} — {', '.join(labels[:5])}{'…' if len(labels)>5 else ''}"

    def _current_preset(self):
        cur = sorted(self._input_exts)
        for name,lst in INPUT_PRESETS.items():
            if sorted(set(x.lower() for x in lst)) == cur:
                return name
        return INPUT_PRESET_CUSTOM

    def _report_hint_text(self):
        d = self.settings.get("report_dir","").strip()
        return f"Сохраняется в: {d}" if d else f"По умолчанию: {app_dir()/'отчет'}"

    def _refresh_section_summaries(self):
        try:
            fmts = [l for fid,l,_ in OUTPUT_FORMATS if self.out_vars.get(fid)]
            ow = "перезапись" if self.settings.get("overwrite_mode") == "overwrite" else "пропуск"
            if not fmts:
                self.sec_formats.set_summary("форматы не выбраны ⚠")
            else:
                s = fmts[0] if len(fmts) == 1 else f"{len(fmts)} формата"
                self.sec_formats.set_summary(f"{s} · {ow}")
            n = len(parse_lines(self.settings.get("articles","")))
            self.sec_articles.set_summary(f"{n} артикулов" if n else "не задано")
            inc = parse_lines(self.settings.get("include_folders",""))
            exc = parse_lines(self.settings.get("exclude_folders",""))
            parts = []
            if inc:
                parts.append(f"Включено папок: {len(inc)}")
            if exc:
                parts.append(f"Исключено папок: {len(exc)}")
            dm = {"all":"всё время","day":"за сутки","week":"за неделю",
                  "month":"за месяц","year":"за год","custom":"свой период"}.get(
                      self.settings.get("date_mode","all"), "")
            if self.settings.get("date_mode","all") != "all":
                parts.append(f"Период: {dm}")
            self.sec_filters.set_summary(" · ".join(parts) if parts else "без фильтра")
            self.sec_stats.set_summary("отчёт вкл" if self.settings.get("save_stats") else "выкл")
        except Exception:
            pass
        try:
            self.page.update()
        except Exception:
            pass

    def _commit_typed_path(self, e=None):
        if self._scanning:
            return
        raw = (self.folder_path.value or "").strip().strip('"').strip("'")
        if not raw:
            return
        p = Path(raw)
        if p.is_dir():
            self.current_folder = p
            self.settings["last_folder"] = str(p)
            save_settings(self.settings)
            self._file_cache = None
            self._start_work(None)
        else:
            self._snack("Такой папки не существует")

    def _set_path_editable(self, v):
        try:
            self.folder_path.read_only = not v
            self.folder_path.update()
        except Exception:
            pass

    def _on_preset_change(self, e):
        name = _ev(e)
        if name == INPUT_PRESET_CUSTOM:
            prev = self._current_preset()
            try:
                self.preset_dd.value = prev
                self.preset_dd.update()
            except Exception:
                pass
            self._open_exts_dialog(None)
            return
        lst = INPUT_PRESETS.get(name)
        if lst:
            self._input_exts = {x.lower() for x in lst}
            self.settings["input_exts"] = sorted(self._input_exts)
            save_settings(self.settings)
            self.exts_summary.value = self._exts_summary()
            self._file_cache = None
            self._refresh_section_summaries()
            self._schedule_rebuild()

    def _build_exts_chips(self, cbs, ext_list, apply_preset, check_all, clear_all):
        active_exts = {ext for ext,cb in zip(ext_list,cbs) if cb.value}
        cur = INPUT_PRESET_CUSTOM
        for name,lst in INPUT_PRESETS.items():
            if sorted(set(x.lower() for x in lst)) == sorted(active_exts):
                cur = name
                break
        chips = []
        for name in INPUT_PRESETS:
            icon_name = PRESET_ICONS.get(name, "CIRCLE")
            chips.append(self._chip(icon_name, name,
                                     on_click=lambda e,n=name: apply_preset(n),
                                     selected=(name == cur)))
        return chips

    def _open_exts_dialog(self, e):
        ext_list = list(INPUT_EXT_INFO.keys())
        cbs = []
        for ext,label in INPUT_EXT_INFO.items():
            cb = _cb(label=f"{ext}   {label}", value=ext in self._input_exts)
            cbs.append(cb)

        state = {"chips_row": None}

        def refresh_chips():
            try:
                state["chips_row"].controls.clear()
                for chip in self._build_exts_chips(cbs, ext_list, apply_preset,
                                                     check_all, clear_all):
                    state["chips_row"].controls.append(chip)
                self.page.update()
            except Exception:
                pass

        def apply_preset(name):
            lst = INPUT_PRESETS.get(name) or []
            for ext,cb in zip(ext_list,cbs):
                cb.value = ext in lst
            refresh_chips()

        def check_all(ev):
            for cb in cbs:
                cb.value = True
            refresh_chips()

        def clear_all(ev):
            for cb in cbs:
                cb.value = False
            refresh_chips()

        def save(ev):
            new = {ext for ext,cb in zip(ext_list,cbs) if cb.value}
            if not new:
                self._snack("Не выбрано ни одного расширения")
                return
            self._input_exts = new
            self.settings["input_exts"] = sorted(new)
            save_settings(self.settings)
            self.exts_summary.value = self._exts_summary()
            self._file_cache = None
            self.preset_dd.value = self._current_preset()
            self._close_dlg(dlg)
            self._refresh_section_summaries()
            self._schedule_rebuild()

        chips_row = ft.Row(self._build_exts_chips(cbs, ext_list, apply_preset,
                                                   check_all, clear_all),
                            spacing=GAP, wrap=True, run_spacing=GAP,
                            run_alignment=ft.MainAxisAlignment.START)
        state["chips_row"] = chips_row

        actions_row = ft.Row([
            self._chip("CHECK_BOX", "Включить все", on_click=check_all),
            self._chip("FILTER_ALT_OFF", "Снять всё", on_click=clear_all),
        ], spacing=GAP, wrap=True)

        cbs_wrapped = []
        for cb in cbs:
            cbs_wrapped.append(ft.Container(content=cb, padding=_pad_sym(0, 6), height=30))

        cbs_list = ft.Column(cbs_wrapped, spacing=0, tight=True)
        cbs_holder = ft.Column([
            cbs_list,
            ft.Divider(height=1),
            ft.Container(height=30),
        ], spacing=GAP, tight=True)

        dlg = ft.AlertDialog(
            modal=True,
            content=ft.Container(
                content=ft.Column([
                    ft.Text("Входные форматы файлов", size=14,
                             weight=ft.FontWeight.BOLD),
                    chips_row,
                    actions_row,
                    ft.Divider(height=1),
                    cbs_holder,
                ], spacing=GAP, tight=True), width=560),
            actions=[
                ft.TextButton("Отмена", on_click=lambda e: self._close_dlg(dlg)),
                ft.FilledButton("Сохранить", on_click=save),
            ])
        self._open_dlg(dlg)

    def _on_output_toggle(self, fid, value):
        self.out_vars[fid] = bool(value)
        self.settings[f"fmt_{fid}"] = bool(value)
        save_settings(self.settings)
        self._refresh_section_summaries()

    def _on_overwrite_change(self, e):
        v = _ev(e,"")
        self.settings["overwrite_mode"] = "overwrite" if v == "Перезаписывать" else "skip"
        save_settings(self.settings)
        self._refresh_section_summaries()
        self._refresh_status_hint()

    def _on_articles_change(self, e):
        self.settings["articles"] = self.articles_text.value or ""
        save_settings(self.settings)
        self._refresh_section_summaries()
        self._schedule_rebuild()

    def _on_match_change(self, e):
        v = _ev(e,"")
        self.settings["match_mode"] = "partial" if v == "Частичное" else "exact"
        save_settings(self.settings)
        self._schedule_rebuild()

    def _on_filters_change(self, e):
        self.settings["include_folders"] = self.include_text.value or ""
        self.settings["exclude_folders"] = self.exclude_text.value or ""
        self.settings["date_from"] = self.date_from_field.value or ""
        self.settings["date_to"] = self.date_to_field.value or ""
        save_settings(self.settings)
        self._refresh_section_summaries()
        self._schedule_rebuild()

    def _on_date_mode_change(self, e):
        v = _ev(e,"Все")
        mp = {"Все":"all","За сутки":"day","За неделю":"week",
              "За месяц":"month","За год":"year","Свой период":"custom"}
        self.settings["date_mode"] = mp.get(v,"all")
        save_settings(self.settings)
        self._refresh_section_summaries()
        self._schedule_rebuild()

    def _on_date_kind_change(self, e):
        v = _ev(e,"")
        self.settings["date_kind"] = "modified" if v == "Изменения файла" else "created"
        save_settings(self.settings)
        self._schedule_rebuild()

    def _on_save_stats(self, value):
        self.settings["save_stats"] = bool(value)
        save_settings(self.settings)
        self._refresh_section_summaries()

    def _on_report_dir_change(self, e):
        self.settings["report_dir"] = self.report_dir_field.value or ""
        save_settings(self.settings)
        self.report_hint.value = self._report_hint_text()

    def _on_preview_bg_change(self, e):
        self.settings["dark_preview"] = bool(self.dark_switch.value)
        save_settings(self.settings)
        self._apply_preview_bg()

    def _apply_preview_bg(self):
        if self.settings.get("dark_preview"):
            self.preview_bg = ft.Colors.BLACK
            try:
                self.preview_hint.color = "#c8c8d8"
            except Exception:
                pass
        else:
            self.preview_bg = C_SURF_HIGH
            try:
                self.preview_hint.color = C_ON_VAR
            except Exception:
                pass
        try:
            self.preview_container.bgcolor = self.preview_bg
            self.preview_container.update()
        except Exception:
            pass

    async def _choose_folder(self, e):
        try:
            path = await self.picker.get_directory_path(dialog_title="Выберите папку")
        except Exception as ex:
            self._snack(f"Диалог: {ex}")
            return
        if not path:
            return
        self.current_folder = Path(path)
        self.folder_path.value = path
        self.settings["last_folder"] = path
        save_settings(self.settings)
        self._file_cache = None
        self._file_meta = {}
        self.processed = set()
        self.finished = False
        try:
            self.folder_path.update()
        except Exception:
            pass
        self._start_work(None)

    def _rescan(self, e):
        if not self.current_folder or not self.current_folder.is_dir():
            self._snack("Сначала выберите папку.")
            return
        self._file_cache = None
        self._file_meta = {}
        self.processed = set()
        self.finished = False
        self._start_work(None)

    def _cancel_scan(self, e):
        self._scan_should_stop = True
        self._snack("Останавливаю сканирование…")

    def _set_scanning(self, v):
        self._scanning = v
        try:
            self.btn_cancel_scan.disabled = not v
            self.btn_rescan.disabled = v
            self.btn_cancel_scan.update()
            self.btn_rescan.update()
        except Exception:
            pass
        self._set_path_editable(not v)

    def _get_psd_list(self, folder, force=False):
        if self._file_cache is not None and not force:
            return self._file_cache
        exts = set(self._input_exts)
        if not exts:
            self._snack("Не выбрано ни одного входного формата.")
            self._file_cache = []
            return []
        self._scan_should_stop = False
        self._set_scanning(True)
        try:
            paths, meta = scan_files(folder, exts,
                                      progress_cb=lambda n,c: self._set_status(f"Сканирую… папок: {n}, файлов: {c}"),
                                      cancel_check=lambda: self._scan_should_stop)
        finally:
            self._set_scanning(False)
        if self._scan_should_stop:
            self._set_status("Сканирование прервано.")
            self._file_cache = None
            self._file_meta = {}
            return None
        seen = set()
        unique = []
        um = {}
        for p in paths:
            try:
                rp = p.resolve()
            except OSError:
                rp = p
            if rp not in seen:
                seen.add(rp)
                unique.append(rp)
                um[rp] = meta.get(p,{"mtime":0.0,"ctime":0.0})
        unique.sort()
        self._file_cache = unique
        self._file_meta = um
        return unique

    def _matches_article(self, path, articles, mode):
        if not articles:
            return True
        s = path.stem.lower()
        if mode == "exact":
            return any(s == a.lower() for a in articles)
        return any(a.lower() in s for a in articles)

    def _matches_folders(self, path, root):
        try:
            rel = path.relative_to(root)
        except ValueError:
            return True
        parts = [p.lower() for p in rel.parts[:-1]]
        inc = [_wildcard(p) for p in parse_lines(self.settings.get("include_folders",""))]
        exc = [_wildcard(p) for p in parse_lines(self.settings.get("exclude_folders",""))]
        for pat in exc:
            for part in parts:
                if fnmatch.fnmatch(part,pat):
                    return False
        if not inc:
            return True
        for pat in inc:
            for part in parts:
                if fnmatch.fnmatch(part,pat):
                    return True
        return False

    def _matches_date(self, path):
        mode = self.settings.get("date_mode","all")
        if mode == "all":
            return True
        meta = self._file_meta.get(path)
        if meta is None:
            try:
                st = path.stat()
                meta = {"mtime":st.st_mtime,"ctime":st.st_ctime}
                self._file_meta[path] = meta
            except OSError:
                return True
        kind = self.settings.get("date_kind","modified")
        ts = meta.get("ctime" if kind == "created" else "mtime",0.0)
        now = datetime.datetime.now().timestamp()
        if mode == "day":
            return ts >= now - 24*3600
        if mode == "week":
            return ts >= now - 7*24*3600
        if mode == "month":
            return ts >= now - 30*24*3600
        if mode == "year":
            return ts >= now - 365*24*3600
        if mode == "custom":
            s = 0.0
            e = float("inf")
            tf = (self.date_from_field.value or "").strip()
            tt = (self.date_to_field.value or "").strip()
            if tf:
                try:
                    s = datetime.datetime.strptime(tf,"%d.%m.%Y").timestamp()
                except Exception:
                    pass
            if tt:
                try:
                    e = (datetime.datetime.strptime(tt,"%d.%m.%Y")+
                          datetime.timedelta(days=1)).timestamp()
                except Exception:
                    pass
            return s <= ts < e
        return True

    def _apply_filter_now(self, e=None):
        self.settings["articles"] = self.articles_text.value or ""
        save_settings(self.settings)
        self._file_cache = None
        self.processed = set()
        self.finished = False
        self._start_work(None)

    def _clear_search(self, e=None):
        self.articles_text.value = ""
        self.settings["articles"] = ""
        save_settings(self.settings)
        self._refresh_section_summaries()
        try:
            self.articles_text.update()
        except Exception:
            pass
        self._schedule_rebuild()

    def _clear_folder_filter(self, e=None):
        self.include_text.value = ""
        self.exclude_text.value = ""
        self.settings["include_folders"] = ""
        self.settings["exclude_folders"] = ""
        save_settings(self.settings)
        self._refresh_section_summaries()
        self._schedule_rebuild()

    def _schedule_rebuild(self):
        if self._debounce_timer is not None:
            try:
                self._debounce_timer.cancel()
            except Exception:
                pass
        self._debounce_timer = threading.Timer(0.7, self._rebuild_now)
        self._debounce_timer.daemon = True
        self._debounce_timer.start()

    def _rebuild_now(self):
        try:
            self._debounce_timer = None
            if self._scanning:
                return
            if not self.current_folder:
                return
            if self._file_cache is None:
                def job():
                    all_files = self._get_psd_list(self.current_folder, force=False)
                    if all_files is None:
                        return
                    self._apply_filters_live()
                _run_thread(self.page, job)
            else:
                _run_thread(self.page, self._apply_filters_live)
        except Exception:
            pass

    def _apply_filters_live(self):
        try:
            folder = self.current_folder
            all_files = self._file_cache or []
            self.all_found_count = len(all_files)
            af = [p for p in all_files if self._matches_folders(p, folder)]
            self.all_after_folder_filter = len(af)
            ad = [p for p in af if self._matches_date(p)]
            self.all_after_date_filter = len(ad)
            arts = parse_lines(self.settings.get("articles",""))
            mode = self.settings.get("match_mode","partial")
            files = [p for p in ad if self._matches_article(p, arts, mode)]
            if not self.finished:
                files = [p for p in files if p not in self.processed]
            self.files = files
            self.idx = 0

            if self.finished:
                if files:
                    self._set_status(
                        f"Найдено {len(files)} из {len(all_files)} файлов. Нажмите СТАРТ (F5).")
                    self._load_current()
                else:
                    self._set_status(
                        f"Найдено {len(all_files)} файлов, под фильтр ничего не подошло.")
                    self.full_image = None
                    self.preview_img.visible = False
                    try:
                        self.preview_img.src = ""
                    except Exception:
                        pass
                    self.preview_hint_box.visible = True
                    self.preview_hint.value = "Под фильтр ничего не подошло"
                    self.name_text.value = ""
                    self.counter_top.value = ""
                    try:
                        self.page.update()
                    except Exception:
                        pass
                return

            if not files:
                self.full_image = None
                self.preview_img.visible = False
                try:
                    self.preview_img.src = ""
                except Exception:
                    pass
                self.preview_hint_box.visible = True
                self.preview_hint.value = ("Все файлы отфильтрованы. "
                                             "Нажмите «Стоп» (Esc) для завершения.")
                self.name_text.value = ""
                self.counter_top.value = "0 / 0"
                try:
                    self.btn_no.disabled = True
                    self.btn_yes.disabled = True
                    self.btn_all.disabled = True
                    self.btn_stop.disabled = False
                except Exception:
                    pass
                try:
                    self.page.update()
                except Exception:
                    pass
                return
            self._load_current()
        except Exception as ex:
            print("live rebuild err:", ex)

    def _scan_folders_dialog(self, e):
        folder = self.current_folder
        if not folder or not folder.is_dir():
            self._snack("Сначала выберите папку.")
            return

        def job():
            all_files = self._get_psd_list(folder)
            if not all_files:
                self._snack("Файлов не найдено.")
                return
            names = set()
            for p in all_files:
                try:
                    rel = p.relative_to(folder)
                except ValueError:
                    continue
                for part in rel.parts[:-1]:
                    names.add(part)
            names = sorted(names, key=str.lower)
            if not names:
                self._snack("Подпапок не найдено.")
                return
            inc_cur = {x.lower() for x in parse_lines(self.settings.get("include_folders",""))}
            exc_cur = {x.lower() for x in parse_lines(self.settings.get("exclude_folders",""))}
            vi, ve = {}, {}
            rows = []
            for name in names:
                a = _cb(value=name.lower() in inc_cur)
                b = _cb(value=name.lower() in exc_cur)
                vi[name], ve[name] = a, b
                rows.append(ft.Container(
                    content=ft.Row([
                        ft.Container(content=a, width=60, alignment=CENTER),
                        ft.Container(content=b, width=70, alignment=CENTER),
                        ft.Text(name, size=TXT_SIZE, expand=True,
                                 overflow=ft.TextOverflow.ELLIPSIS),
                    ], spacing=GAP, vertical_alignment=ft.CrossAxisAlignment.CENTER),
                    padding=_pad_sym(PAD, 4),
                    border_radius=6,
                    bgcolor=C_SURF_HIGHEST,
                ))

            def ok(ev):
                inc = [n for n in names if vi[n].value]
                exc = [n for n in names if ve[n].value]
                both = set(inc) & set(exc)
                if both:
                    self._snack("Одна папка и в «Искать», и в «Пропуск»: " + ", ".join(sorted(both)))
                    return
                self.include_text.value = "\n".join(inc)
                self.exclude_text.value = "\n".join(exc)
                self.settings["include_folders"] = self.include_text.value
                self.settings["exclude_folders"] = self.exclude_text.value
                save_settings(self.settings)
                self._close_dlg(dlg)
                self._refresh_section_summaries()
                try:
                    self.include_text.update()
                    self.exclude_text.update()
                except Exception:
                    pass
                self._schedule_rebuild()

            def check_all_incl(ev):
                for cb in vi.values():
                    cb.value = True
                self.page.update()

            def check_all_excl(ev):
                for cb in ve.values():
                    cb.value = True
                self.page.update()

            def clear_all(ev):
                for cb in vi.values():
                    cb.value = False
                for cb in ve.values():
                    cb.value = False
                self.page.update()

            btn_left_col = ft.Column([
                ft.OutlinedButton("Искать все", on_click=check_all_incl,
                                   height=ROW_H, expand=True,
                                   style=_btn_style()),
                ft.OutlinedButton("Отмена", on_click=lambda ev: self._close_dlg(dlg),
                                   height=ROW_H, expand=True,
                                   style=_btn_style()),
            ], spacing=GAP, expand=True, horizontal_alignment=ft.CrossAxisAlignment.STRETCH)
            btn_right_col = ft.Column([
                ft.OutlinedButton("Пропускать все", on_click=check_all_excl,
                                   height=ROW_H, expand=True,
                                   style=_btn_style()),
                ft.FilledButton("ОК", on_click=ok,
                                 height=ROW_H, expand=True,
                                 style=_btn_style()),
            ], spacing=GAP, expand=True, horizontal_alignment=ft.CrossAxisAlignment.STRETCH)

            buttons_row = ft.Row([btn_left_col, btn_right_col], spacing=GAP)

            clear_row = ft.Row([
                ft.OutlinedButton("Снять все отметки", on_click=clear_all,
                                   height=ROW_H, style=_btn_style()),
                ft.Container(expand=True),
                ft.Text(f"Всего папок: {len(names)}", size=HINT_SIZE, color=C_ON_VAR, italic=True),
            ], spacing=GAP, vertical_alignment=ft.CrossAxisAlignment.CENTER)

            list_box = ft.Container(
                content=ft.Column(rows, spacing=4, tight=True, scroll=ft.ScrollMode.AUTO),
                height=340, padding=6, border_radius=8,
                bgcolor=C_SURF_HIGH,
                border=_border_all(1, C_OUT_VAR),
            )

            header_row = ft.Row([
                ft.Container(content=ft.Text("Искать", size=LBL_SIZE, weight=ft.FontWeight.BOLD),
                              width=60, alignment=CENTER),
                ft.Container(content=ft.Text("Пропуск", size=LBL_SIZE, weight=ft.FontWeight.BOLD),
                              width=70, alignment=CENTER),
                ft.Text("Папка", size=LBL_SIZE, weight=ft.FontWeight.BOLD, expand=True),
            ], spacing=GAP)

            body_col = ft.Column([
                ft.Text("Отметьте «Искать» для нужных папок и «Пропуск» — для ненужных.",
                        size=HINT_SIZE, color=C_ON_VAR, italic=True),
                ft.Divider(height=1),
                header_row,
                list_box,
                clear_row,
                ft.Divider(height=1),
                buttons_row,
            ], spacing=GAP, tight=True)

            dlg = ft.AlertDialog(
                modal=True,
                title=ft.Text("Какие папки искать, а какие — пропускать", size=14,
                               weight=ft.FontWeight.BOLD),
                content=ft.Container(content=body_col, width=600),
            )
            self._open_dlg(dlg)
        _run_thread(self.page, job)

    async def _load_articles_from_file(self, e):
        try:
            files = await self.picker.pick_files(
                dialog_title="Файл со списком артикулов",
                allowed_extensions=["txt","csv","list","xlsx","xlsm","docx"],
                allow_multiple=False)
        except Exception as ex:
            self._snack(f"Диалог: {ex}")
            return
        if not files:
            return
        path = Path(files[0].path)
        try:
            if path.suffix.lower() in (".xlsx",".xlsm"):
                lines = self._load_from_xlsx(path)
            elif path.suffix.lower() == ".docx":
                lines = self._clean_lines(read_docx(path))
            else:
                lines = self._clean_lines(self._read_text_file(path))
            if not lines:
                self._snack("Не найдено артикулов.")
                return
            self._apply_loaded_articles(lines, path.name)
        except Exception as ex:
            self._snack(f"Ошибка: {ex}")

    @staticmethod
    def _read_text_file(p):
        for enc in ("utf-8-sig","utf-8","cp1251","latin-1"):
            try:
                return p.read_text(encoding=enc)
            except UnicodeDecodeError:
                continue
        raise ValueError("Кодировка не определена")

    @staticmethod
    def _clean_lines(raw):
        if isinstance(raw,str):
            raw = raw.splitlines()
        return [str(x).strip().lstrip("\ufeff") for x in raw if str(x).strip()]

    def _load_from_xlsx(self, path):
        try:
            order, data = read_xlsx(path)
        except RuntimeError as e:
            self._snack(str(e))
            return []
        if not order:
            return []
        if len(order) == 1:
            return data[order[0]]
        chosen = self._ask_xlsx_column_dialog(order, data)
        if not chosen:
            return []
        self.settings["last_xlsx_column"] = chosen
        save_settings(self.settings)
        return data[chosen]

    def _ask_xlsx_column_dialog(self, order, data):
        result = {"col":None}
        event = threading.Event()
        default = self.settings.get("last_xlsx_column","")
        rows = []
        for col in order:
            vals = data[col]
            prev = ", ".join(vals[:5])
            if len(vals) > 5:
                prev += f", … (+{len(vals)-5})"
            rd = ft.Radio(value=col, label="")
            rows.append(ft.Row([
                ft.Container(content=rd, width=30),
                ft.Text(f"Столбец {col} ({len(vals)}):", size=LBL_SIZE, weight=ft.FontWeight.BOLD),
                ft.Text(prev, size=HINT_SIZE, color=C_ON_VAR, expand=True),
            ], spacing=GAP))
        rg = ft.RadioGroup(content=ft.Column(rows, spacing=4,
                                              scroll=ft.ScrollMode.AUTO, height=400),
                            value=default if default in order else order[0])

        def ok(ev):
            result["col"] = rg.value
            self._close_dlg(dlg)
            event.set()

        def cancel(ev):
            result["col"] = None
            self._close_dlg(dlg)
            event.set()

        dlg = ft.AlertDialog(
            modal=True,
            title=ft.Text("Выбор столбца"),
            content=ft.Container(content=ft.Column([
                ft.Text("Первые 5 непустых значений каждого столбца.", size=HINT_SIZE),
                rg,
            ], spacing=GAP, tight=True), width=620),
            actions=[ft.TextButton("Отмена", on_click=cancel),
                     ft.FilledButton("Загрузить", on_click=ok)])
        self._open_dlg(dlg)
        event.wait(timeout=120)
        return result["col"]

    def _apply_loaded_articles(self, lines, source):
        self.articles_text.value = "\n".join(lines)
        self.settings["articles"] = self.articles_text.value
        save_settings(self.settings)
        self._snack(f"Загружено {len(lines)} из {source}")
        self._refresh_section_summaries()
        try:
            self.articles_text.update()
        except Exception:
            pass
        self._schedule_rebuild()

    def _load_articles_from_gsheet(self, e):
        url_tf = _tf(value=self.settings.get("gsheet_url",""), hint_text="Ссылка")
        method_dd = _dd(value="Публичная таблица" if self.settings.get("gsheet_method","public") == "public" else "Service Account",
                         options=[ft.dropdown.Option("Публичная таблица"),
                                  ft.dropdown.Option("Service Account")])
        creds_tf = _tf(value=self.settings.get("gsheet_creds",""), hint_text="JSON-ключ")

        def do(ev):
            url = url_tf.value.strip()
            if not url:
                self._snack("Вставьте ссылку.")
                return
            method = "public" if method_dd.value == "Публичная таблица" else "private"
            self.settings["gsheet_url"] = url
            self.settings["gsheet_method"] = method
            self.settings["gsheet_creds"] = creds_tf.value.strip()
            save_settings(self.settings)
            self._close_dlg(dlg)

            def job():
                try:
                    if method == "public":
                        order,data = read_gsheet_public(url)
                    else:
                        order,data = read_gsheet_private(url, creds_tf.value.strip())
                except Exception as ex:
                    self._snack(f"Ошибка: {ex}")
                    return
                if not order:
                    self._snack("Нет данных.")
                    return
                if len(order) == 1:
                    chosen = order[0]
                else:
                    chosen = self._ask_xlsx_column_dialog(order, data)
                    if not chosen:
                        return
                self.settings["gsheet_column"] = chosen
                save_settings(self.settings)
                lines = data.get(chosen, [])
                if not lines:
                    self._snack("Столбец пуст.")
                    return
                self._apply_loaded_articles(lines, f"Sheets ({chosen})")
            _run_thread(self.page, job)

        dlg = ft.AlertDialog(
            modal=True,
            title=ft.Text("Загрузить из Google Sheets"),
            content=ft.Container(content=ft.Column([
                self._field("Ссылка:", url_tf),
                self._spacer(),
                self._field("Метод доступа:", method_dd),
                self._spacer(),
                self._field("JSON-ключ (для Service Account):", creds_tf,
                            hint="Для публичной таблицы не нужно"),
            ], spacing=GAP_FIELD, tight=True), width=560),
            actions=[ft.TextButton("Отмена", on_click=lambda e: self._close_dlg(dlg)),
                     ft.FilledButton("Загрузить", on_click=do)])
        self._open_dlg(dlg)

    def _choose_report_dir(self, e):
        async def job():
            try:
                path = await self.picker.get_directory_path(dialog_title="Папка отчётов")
            except Exception as ex:
                self._snack(f"Диалог: {ex}")
                return
            if not path:
                return
            self.report_dir_field.value = path
            self.settings["report_dir"] = path
            save_settings(self.settings)
            self.report_hint.value = self._report_hint_text()
            try:
                self.report_dir_field.update()
                self.report_hint.update()
            except Exception:
                pass
        try:
            self.page.run_task(job)
        except Exception:
            _run_thread(self.page, lambda: None)

    def _reset_report_dir(self, e):
        self.report_dir_field.value = ""
        self.settings["report_dir"] = ""
        save_settings(self.settings)
        self.report_hint.value = self._report_hint_text()
        try:
            self.report_dir_field.update()
            self.report_hint.update()
        except Exception:
            pass

    def _report_dir(self):
        c = self.settings.get("report_dir","").strip()
        return Path(c) if c else app_dir()/"отчет"

    def _write_stats_file(self):
        now = datetime.datetime.now()
        try:
            d = self._report_dir()
            d.mkdir(parents=True,exist_ok=True)
        except Exception:
            d = self.current_folder or app_dir()
        out = d / f"_pixforge_stats_{now:%Y-%m-%d_%H-%M-%S}.txt"
        L = ["="*70,f"{APP_NAME} — отчёт","="*70,
             f"Версия: {APP_VERSION}",
             f"Дата и время: {now:%Y-%m-%d %H:%M:%S}",
             f"Папка обработки: {self.current_folder}",
             f"Папка отчёта: {d}","",
             "Настройки:",
             f"  Входные: {', '.join(sorted(self._input_exts))}",
             f"  Выходные: {', '.join(l for fid,l,_ in OUTPUT_FORMATS if self.out_vars.get(fid)) or '(нет)'}",
             f"  Режим: {'перезапись' if self.settings.get('overwrite_mode') == 'overwrite' else 'пропуск'}",
             f"  Совпадение: {'полное' if self.settings.get('match_mode') == 'exact' else 'частичное'}",
             f"  Артикулов: {len(parse_lines(self.settings.get('articles','')))}",
             f"  Включ. папки: {', '.join(parse_lines(self.settings.get('include_folders',''))) or '(все)'}",
             f"  Исключ. папки: {', '.join(parse_lines(self.settings.get('exclude_folders',''))) or '(нет)'}",
             f"  Фильтр по дате: {self.settings.get('date_mode')}",
             "", "Итоги:",
             f"  Всего файлов: {self.all_found_count}",
             f"  После папок: {self.all_after_folder_filter}",
             f"  После даты: {self.all_after_date_filter}",
             f"  Подошло: {len(self.files)}",
             f"  Сконвертировано: {self.converted}",
             f"  Пропущено: {self.skipped}",
             f"  Ошибок: {self.errors}",
             "", "="*70, "Детали:", "="*70]
        for e in self.log_entries:
            p = e["path"]
            try:
                rel = p.relative_to(self.current_folder)
            except Exception:
                rel = p
            a = e["action"]
            if a == "converted":
                w = ", ".join(e.get("written",[])) or "—"
                L.append(f"[OK]    {rel}   →   {w}")
            elif a == "skipped_existing":
                L.append(f"[SKIP]  {rel}   уже есть")
            elif a == "skipped_manual":
                L.append(f"[SKIP]  {rel}   вручную")
            elif a == "error":
                L.append(f"[ERR]   {rel}   {e.get('error','')}")
        out.write_text("\n".join(L),encoding="utf-8")
        return out

    def _set_status(self, text):
        self.status_text.value = text
        try:
            self.status_text.update()
        except Exception:
            pass

    def _refresh_status_hint(self):
        if self.finished or self.idx >= len(self.files):
            return
        path = self.files[self.idx]
        fmts = self._format_names()
        hint = "Форматы: " + (", ".join(fmts) if fmts else "не выбраны ⚠")
        ex = self._existing_outputs(path)
        if ex:
            if self.settings.get("overwrite_mode") == "skip":
                hint += f"  ·  пропустится: {', '.join(ex)}"
            else:
                hint += f"  ·  перезапишется: {', '.join(ex)}"
        else:
            hint += "  ·  Y / → сохранить, N / ← пропустить"
        self._set_status(hint)

    def _start_work(self, e=None):
        if not self.current_folder or not self.current_folder.is_dir():
            self._snack("Сначала выберите папку.")
            return
        miss = []
        for fid,(imp,pip) in OUTPUT_REQUIRED_LIB.items():
            if self.out_vars.get(fid) and not _is_installed(imp):
                lbl = next((l for f,l,_ in OUTPUT_FORMATS if f == fid),fid)
                miss.append((lbl,pip))
        if miss:
            self._snack("Не хватает: " + ", ".join(m[0] for m in miss))
            return
        self.finished = False
        self.processed = set()

        def job():
            folder = self.current_folder
            all_files = self._get_psd_list(folder, force=True)
            if all_files is None:
                return
            self.all_found_count = len(all_files)
            af = [p for p in all_files if self._matches_folders(p,folder)]
            self.all_after_folder_filter = len(af)
            ad = [p for p in af if self._matches_date(p)]
            self.all_after_date_filter = len(ad)
            arts = parse_lines(self.settings.get("articles",""))
            mode = self.settings.get("match_mode","partial")
            files = [p for p in ad if self._matches_article(p,arts,mode)]
            self.files = files
            self.idx = 0
            self.converted = 0
            self.skipped = 0
            self.errors = 0
            self.log_entries = []
            if not files:
                self._set_status(f"Найдено {len(all_files)}, но по фильтрам ничего не подошло.")
                self.finished = True
                return
            self.start_box.visible = False
            self._set_actions(True)
            try:
                self.start_box.update()
            except Exception:
                pass
            self._load_current()
        _run_thread(self.page, job)

    def _set_actions(self, v):
        for b in (self.btn_yes,self.btn_no,self.btn_all,self.btn_stop):
            b.disabled = not v
        try:
            self.page.update()
        except Exception:
            pass

    def _update_counter(self):
        total = len(self.files)
        done = len(self.processed)
        if total:
            base = f"{self.idx+1} / {total}"
            if done:
                base += f"  ·  обработано: {done}"
        else:
            base = "0 / 0"
        self.counter_top.value = base
        try:
            self.counter_top.update()
        except Exception:
            pass

    def _load_current(self):
        self.full_image = None
        self.preview_img.visible = False
        try:
            self.preview_img.src = ""
        except Exception:
            pass
        if self.idx >= len(self.files):
            self.preview_img.visible = False
            self.preview_hint_box.visible = True
            self.preview_hint.value = ("Обработка завершена. "
                                         "Нажмите «Стоп / Итог» (Esc), чтобы увидеть итоги.")
            self.name_text.value = ""
            self.counter_top.value = (f"{self.idx} / {self.idx}" if self.idx else "0 / 0")
            try:
                self.btn_no.disabled = True
                self.btn_yes.disabled = True
                self.btn_all.disabled = True
                self.btn_stop.disabled = False
            except Exception:
                pass
            try:
                self.page.update()
            except Exception:
                pass
            return
        path = self.files[self.idx]
        self._update_counter()
        try:
            rel = path.relative_to(self.current_folder)
        except Exception:
            rel = path
        self.name_text.value = f"{path.name}  —  {rel.parent}"
        if not self.finished:
            self._set_status("Загрузка превью…")
        try:
            self.page.update()
        except Exception:
            pass

        def job():
            try:
                img = load_image_any(path)
                self.full_image = img
                data_uri = pil_to_data_uri(img)
                self.preview_img.src = data_uri
                self.preview_img.visible = True
                self.preview_hint_box.visible = False
                if not self.finished:
                    self._refresh_status_hint()
            except Exception as ex:
                self.preview_hint_box.visible = True
                self.preview_hint.value = f"⚠ Не удалось открыть файл:\n{ex}"
                if not self.finished:
                    self._set_status("Ошибка чтения. N / ← — пропустить.")
            try:
                self.page.update()
            except Exception:
                pass
        _run_thread(self.page, job)

    def _format_names(self):
        return [l for fid,l,_ in OUTPUT_FORMATS if self.out_vars.get(fid)]

    def _target_paths(self, path):
        stem, parent = path.stem, path.parent
        wt = self.out_vars["png_t"]
        out = []
        for fid,_,_ in OUTPUT_FORMATS:
            if not self.out_vars[fid]:
                continue
            if fid == "png_t":
                out.append((parent/f"{stem}.png","png_t"))
            elif fid == "png_o":
                name = f"{stem}_white.png" if wt else f"{stem}.png"
                out.append((parent/name,"png_o"))
            elif fid == "jpg":
                out.append((parent/f"{stem}.jpg","jpg"))
            elif fid == "webp":
                out.append((parent/f"{stem}.webp","webp"))
            elif fid == "tiff":
                out.append((parent/f"{stem}.tiff","tiff"))
            elif fid == "bmp":
                out.append((parent/f"{stem}.bmp","bmp"))
            elif fid == "pdf":
                out.append((parent/f"{stem}.pdf","pdf"))
            elif fid == "avif":
                out.append((parent/f"{stem}.avif","avif"))
        return out

    def _existing_outputs(self, path):
        return [p.name for p,_ in self._target_paths(path) if p.exists()]

    def _save_image(self, img, path):
        targets = self._target_paths(path)
        if not targets:
            raise ValueError("Не выбран формат вывода")
        ow = self.settings.get("overwrite_mode") == "overwrite"
        w, sk = [], []
        for out,kind in targets:
            if out.exists() and not ow:
                sk.append(out.name)
                continue
            save_image_in_format(img, out, kind)
            w.append(out.name)
        return w, sk

    def _accept(self, e=None):
        if self.finished or self.idx >= len(self.files):
            return
        path = self.files[self.idx]
        try:
            if self.full_image is None:
                raise ValueError("Изображение не загружено")
            w, sk = self._save_image(self.full_image, path)
            if w:
                self.converted += 1
                self.log_entries.append({"path":path,"action":"converted","written":w,"skipped":sk})
            elif sk:
                self.skipped += 1
                self.log_entries.append({"path":path,"action":"skipped_existing","written":[],"skipped":sk})
        except Exception as ex:
            self.errors += 1
            self.log_entries.append({"path":path,"action":"error","error":str(ex)})
            self._snack(f"Ошибка: {ex}")
        self.processed.add(path)
        self.idx += 1
        self._load_current()

    def _reject(self, e=None):
        if self.finished or self.idx >= len(self.files):
            return
        path = self.files[self.idx]
        self.skipped += 1
        self.log_entries.append({"path":path,"action":"skipped_manual"})
        self.processed.add(path)
        self.idx += 1
        self._load_current()

    def _accept_all(self, e=None):
        if self.finished:
            return

        def job():
            while self.idx < len(self.files):
                path = self.files[self.idx]
                try:
                    img = self.full_image or load_image_any(path)
                    w, sk = self._save_image(img, path)
                    if w:
                        self.converted += 1
                        self.log_entries.append({"path":path,"action":"converted","written":w,"skipped":sk})
                    elif sk:
                        self.skipped += 1
                        self.log_entries.append({"path":path,"action":"skipped_existing","written":[],"skipped":sk})
                except Exception as ex:
                    self.errors += 1
                    self.log_entries.append({"path":path,"action":"error","error":str(ex)})
                self.processed.add(path)
                self.full_image = None
                self.idx += 1
            self.preview_img.visible = False
            self.preview_hint_box.visible = True
            self.preview_hint.value = ("Обработка завершена. "
                                         "Нажмите «Стоп / Итог» (Esc), чтобы увидеть итоги.")
            self.name_text.value = ""
            self.counter_top.value = f"{self.idx} / {self.idx}"
            try:
                self.btn_no.disabled = True
                self.btn_yes.disabled = True
                self.btn_all.disabled = True
                self.btn_stop.disabled = False
            except Exception:
                pass
            try:
                self.page.update()
            except Exception:
                pass
        _run_thread(self.page, job)

    def _finish(self, e=None):
        if self.finished:
            return
        if e is None and not self.log_entries:
            return
        self.finished = True
        stats = None
        if self.settings.get("save_stats"):
            try:
                stats = self._write_stats_file()
            except Exception:
                pass
        msg = (f"Обработка завершена.\n\n"
               f"Всего файлов: {self.all_found_count}\n"
               f"После папок: {self.all_after_folder_filter}\n"
               f"После даты: {self.all_after_date_filter}\n"
               f"Подошло под артикулы: {len(self.files)}\n"
               f"Сконвертировано: {self.converted}\n"
               f"Пропущено: {self.skipped}\n"
               f"Ошибок: {self.errors}")
        if stats:
            msg += f"\n\nОтчёт: {stats}"

        def close_dlg(ev):
            self._close_dlg(dlg)

        dlg = ft.AlertDialog(
            modal=True,
            title=ft.Text("Итог"),
            content=ft.Text(msg),
            actions=[ft.FilledButton("OK", autofocus=True,
                                       on_click=close_dlg)])
        self._open_dlg(dlg)
        self.start_box.visible = True
        self.preview_img.visible = False
        self.preview_hint_box.visible = True
        self.preview_hint.value = "Готово. Нажмите СТАРТ (F5)."
        self._set_actions(False)
        self.counter_top.value = ""
        self.name_text.value = ""
        self.status_text.value = ""
        try:
            self.page.update()
        except Exception:
            pass

    def _on_key(self, e):
        try:
            key = (e.key or "").lower()
        except Exception:
            return
        if key == "escape" and self._active_dialog is not None:
            try:
                self._close_dlg(self._active_dialog)
            except Exception:
                pass
            return
        if key == "f5":
            self._start_work(None)
            return
        if self.finished:
            return
        if key in ("n","arrow left"):
            self._reject()
        elif key in ("y","arrow right"):
            self._accept()
        elif key == "a":
            self._accept_all()
        elif key == "escape":
            self._finish()

# ============================================================================
#  ТОЧКА ВХОДА
# ============================================================================
if __name__ == "__main__":
    try:
        ft.app(target=ConverterApp)
    except Exception as e:
        tb = traceback.format_exc()
        try:
            LOG_PATH.write_text(tb, encoding="utf-8")
        except Exception:
            pass
        print(f"Ошибка: {e}\n\n{tb}")
        input("Enter...")